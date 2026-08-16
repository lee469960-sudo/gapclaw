"""Build online preview payloads for common office / text file types."""

from __future__ import annotations

import csv
import html
import io
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

MAX_EXCEL_ROWS = 500
MAX_EXCEL_COLS = 50
MAX_TEXT_CHARS = 200_000
MAX_PPT_SLIDES = 80
MAX_CSV_ROWS = 500
MAX_CSV_COLS = 50

EXCEL_EXTS = {".xlsx", ".xlsm"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".svg"}
TEXT_EXTS = {
    ".txt",
    ".md",
    ".markdown",
    ".json",
    ".xml",
    ".yaml",
    ".yml",
    ".log",
    ".csv",
    ".tsv",
    ".py",
    ".js",
    ".ts",
    ".vue",
    ".html",
    ".css",
    ".sql",
    ".sh",
    ".ini",
    ".conf",
    ".toml",
}


def _cell_value(cell) -> str:
    if cell is None:
        return ""
    return str(cell)


def _normalize_sheet_rows(rows: list[list]) -> list[list[str]]:
    if not rows:
        return [[""]]
    max_cols = max(len(row) for row in rows)
    return [
        [_cell_value(cell) for cell in row] + [""] * (max_cols - len(row))
        for row in rows
    ]


def read_excel_preview(target: Path) -> dict:
    import openpyxl

    wb = openpyxl.load_workbook(target, read_only=True, data_only=True)
    sheets = []
    truncated = False
    try:
        for name in wb.sheetnames:
            ws = wb[name]
            rows = []
            for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
                if row_idx >= MAX_EXCEL_ROWS:
                    truncated = True
                    break
                cells = [_cell_value(cell) for cell in row[:MAX_EXCEL_COLS]]
                if len(row) > MAX_EXCEL_COLS:
                    truncated = True
                rows.append(cells)
            sheets.append({"name": name, "rows": _normalize_sheet_rows(rows)})
    finally:
        wb.close()
    return {"ok": True, "type": "excel", "sheets": sheets, "truncated": truncated}


def _docx_to_html(target: Path) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(target))
    parts: list[str] = []

    def emit_paragraph(p: Paragraph) -> None:
        style = (p.style.name if p.style is not None else "") or ""
        text = html.escape(p.text or "")
        if not text.strip():
            parts.append("<p><br/></p>")
            return
        if style.startswith("Heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "2"
            level = str(min(max(int(level), 1), 6))
            parts.append(f"<h{level}>{text}</h{level}>")
        else:
            parts.append(f"<p>{text}</p>")

    def emit_table(t: Table) -> None:
        parts.append('<table class="docx-table">')
        for i, row in enumerate(t.rows):
            parts.append("<tr>")
            tag = "th" if i == 0 else "td"
            for cell in row.cells:
                parts.append(f"<{tag}>{html.escape(cell.text or '')}</{tag}>")
            parts.append("</tr>")
        parts.append("</table>")

    for block in doc.element.body:
        tag = block.tag.split("}")[-1] if "}" in block.tag else block.tag
        if tag == "p":
            emit_paragraph(Paragraph(block, doc))
        elif tag == "tbl":
            emit_table(Table(block, doc))

    body = "\n".join(parts) or "<p>（空文档）</p>"
    return body[:MAX_TEXT_CHARS]


def _pptx_slides(target: Path) -> dict:
    """Extract slide text from pptx without requiring python-pptx."""
    slides: list[dict] = []
    truncated = False
    ns = {
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    }
    with zipfile.ZipFile(target, "r") as zf:
        names = sorted(
            n for n in zf.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")
        )
        if len(names) > MAX_PPT_SLIDES:
            truncated = True
            names = names[:MAX_PPT_SLIDES]
        for idx, name in enumerate(names, 1):
            root = ET.fromstring(zf.read(name))
            texts = [
                (node.text or "").strip()
                for node in root.findall(".//a:t", ns)
                if (node.text or "").strip()
            ]
            # de-dupe consecutive repeats common in pptx runs
            compact: list[str] = []
            for t in texts:
                if not compact or compact[-1] != t:
                    compact.append(t)
            title = compact[0] if compact else f"幻灯片 {idx}"
            slides.append({"index": idx, "title": title, "texts": compact})
    return {"ok": True, "type": "pptx", "slides": slides, "truncated": truncated}


def _csv_preview(target: Path) -> dict:
    raw = target.read_text(encoding="utf-8", errors="replace")
    dialect = csv.excel
    sample = raw[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
    except csv.Error:
        pass
    reader = csv.reader(io.StringIO(raw), dialect)
    rows: list[list[str]] = []
    truncated = False
    for i, row in enumerate(reader):
        if i >= MAX_CSV_ROWS:
            truncated = True
            break
        cells = [_cell_value(c) for c in row[:MAX_CSV_COLS]]
        if len(row) > MAX_CSV_COLS:
            truncated = True
        rows.append(cells)
    return {
        "ok": True,
        "type": "excel",
        "sheets": [{"name": target.name, "rows": _normalize_sheet_rows(rows)}],
        "truncated": truncated,
    }


def preview_file(target: Path) -> dict:
    """Return a JSON-serializable preview payload for *target*."""
    if not target.is_file():
        return {"ok": False, "msg": "不是文件"}

    ext = target.suffix.lower()

    if ext == ".pdf":
        return {"ok": True, "type": "pdf"}
    if ext in IMAGE_EXTS:
        mime = {
            ".png": "image/png",
            ".jpg": "image/jpeg",
            ".jpeg": "image/jpeg",
            ".gif": "image/gif",
            ".webp": "image/webp",
            ".bmp": "image/bmp",
            ".svg": "image/svg+xml",
        }.get(ext, "application/octet-stream")
        return {"ok": True, "type": "image", "mime": mime}
    if ext in EXCEL_EXTS:
        try:
            return read_excel_preview(target)
        except Exception as exc:
            return {"ok": False, "msg": f"无法读取 Excel：{exc}"}
    if ext == ".xls":
        return {
            "ok": False,
            "type": "unsupported",
            "msg": "旧版 .xls 请下载后用 Excel 打开；建议另存为 .xlsx",
        }
    if ext == ".docx":
        try:
            return {"ok": True, "type": "docx", "html": _docx_to_html(target)}
        except Exception as exc:
            return {"ok": False, "msg": f"无法读取 Word：{exc}"}
    if ext == ".doc":
        return {
            "ok": False,
            "type": "unsupported",
            "msg": "旧版 .doc 请转换为 .docx 后预览",
        }
    if ext == ".pptx":
        try:
            return _pptx_slides(target)
        except Exception as exc:
            return {"ok": False, "msg": f"无法读取 PPT：{exc}"}
    if ext == ".ppt":
        return {
            "ok": False,
            "type": "unsupported",
            "msg": "旧版 .ppt 请转换为 .pptx 后预览",
        }
    if ext in {".csv", ".tsv"}:
        try:
            return _csv_preview(target)
        except Exception as exc:
            return {"ok": False, "msg": f"无法读取 CSV：{exc}"}
    if ext in {".md", ".markdown"}:
        content = target.read_text(encoding="utf-8", errors="replace")[:MAX_TEXT_CHARS]
        return {"ok": True, "type": "markdown", "content": content}

    # default: try text
    try:
        content = target.read_text(encoding="utf-8", errors="replace")[:MAX_TEXT_CHARS]
        # Heuristic: mostly binary → refuse
        if "\x00" in content[:2000]:
            return {
                "ok": False,
                "type": "unsupported",
                "msg": f"暂不支持在线预览「{ext or '未知'}」格式，请下载后查看",
            }
        return {"ok": True, "type": "text", "content": content}
    except Exception:
        return {
            "ok": False,
            "type": "unsupported",
            "msg": f"暂不支持在线预览「{ext or '未知'}」格式，请下载后查看",
        }
