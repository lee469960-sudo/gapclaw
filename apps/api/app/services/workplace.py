import json
import os
import shutil
import zipfile
from datetime import datetime
from pathlib import Path

from app.config import get_settings

EXCEL_EXTS = {".xlsx", ".xlsm"}
MAX_EXCEL_ROWS = 500
MAX_EXCEL_COLS = 50


def _benchmark_workplace_root() -> Path | None:
    raw = (os.environ.get("REACT_BENCH_WORKSPACE_ROOT") or "").strip()
    if not raw:
        return None
    return Path(raw).resolve()


def _wp_root(sandbox_id: str) -> Path:
    bench_root = _benchmark_workplace_root()
    if bench_root is not None:
        return bench_root
    sid = (sandbox_id or "").strip() or "default"
    return Path(get_settings().workplace_dir) / sid / "workplace"


def _resolved_wp_root(sandbox_id: str) -> Path:
    return _wp_root(sandbox_id).resolve()


def workplace_root(sandbox_id: str) -> Path:
    return _wp_root(sandbox_id)


def benchmark_workplace_root() -> Path | None:
    return _benchmark_workplace_root()


def _safe_filename(filename: str) -> str:
    if not filename:
        return ""
    name = filename.replace("\\", "/").split("/")[-1].strip()
    if not name or name in (".", ".."):
        return ""
    return name


def _path_under_root(root: Path, rel: str = "") -> Path | None:
    rel = (rel or "").strip().strip("/")
    try:
        target = root if not rel else (root / rel).resolve()
    except (OSError, ValueError):
        return None
    root_s = str(root)
    target_s = str(target)
    if target_s == root_s or target_s.startswith(root_s + os.sep):
        return target
    return None


def ensure_workplace(sandbox_id: str) -> Path:
    """Ensure workplace root and task scratch directory exist."""
    root = _wp_root(sandbox_id)
    root.mkdir(parents=True, exist_ok=True)
    (root / "task").mkdir(exist_ok=True)
    return root




def list_dir(sandbox_id: str, rel: str = "") -> list[dict]:
    root = ensure_workplace(sandbox_id)
    clean = (rel or "").strip().strip("/")
    # Never auto-create paths from shell noise like "nodes/ 2>&1 | head -50"
    if any(ch in clean for ch in "|;&<>`$()"):
        return []
    target = root / clean if clean else root
    if not target.exists() or not target.is_dir():
        return []
    entries = []
    for item in sorted(target.iterdir(), key=lambda x: (not x.is_dir(), x.name.lower())):
        st = item.stat()
        entries.append({
            "name": item.name,
            "is_dir": item.is_dir(),
            "size": st.st_size if item.is_file() else 0,
            "modified": datetime.fromtimestamp(st.st_mtime).strftime("%Y-%m-%d %H:%M"),
            "path": str((Path(clean) / item.name).as_posix()).lstrip("/") if clean else item.name,
        })
    return entries


def list_dirs(sandbox_id: str) -> list[str]:
    root = _wp_root(sandbox_id)
    dirs = []
    for p in root.rglob("*"):
        if p.is_dir():
            dirs.append(str(p.relative_to(root)))
    return dirs


def format_dir_listing(sandbox_id: str, rel: str = "") -> str:
    """Human-readable workplace listing (same data source as UI sidebar)."""
    entries = list_dir(sandbox_id, rel)
    if not entries:
        prefix = f"/workplace/{rel}".rstrip("/") if rel else "/workplace"
        return f"{prefix}\n(空目录)"
    lines = []
    prefix = f"/workplace/{rel}".rstrip("/") if rel else "/workplace"
    lines.append(prefix + ":")
    for item in entries:
        kind = "d" if item["is_dir"] else "-"
        size = item["size"] if not item["is_dir"] else 0
        lines.append(f"{kind} {item['name']}\t{size}\t{item.get('modified', '')}")
    return "\n".join(lines)


def is_valid_deliverable_file(path: Path) -> bool:
    """Non-empty data file; xlsx/xlsm must look like a real zip workbook."""
    if not path or not path.is_file():
        return False
    try:
        size = path.stat().st_size
    except OSError:
        return False
    if size <= 0:
        return False
    suf = path.suffix.lower()
    if suf in {".xlsx", ".xlsm"}:
        try:
            with path.open("rb") as fh:
                magic = fh.read(2)
        except OSError:
            return False
        if magic != b"PK":
            return False
        try:
            return zipfile.is_zipfile(path)
        except OSError:
            return False
    if suf == ".xls":
        return size >= 512
    if suf == ".csv":
        return True
    return True




def find_files(sandbox_id: str, name: str) -> list[str]:
    """Find files by basename under workplace (recursive)."""
    root = ensure_workplace(sandbox_id)
    target = name.strip()
    if not target:
        return []
    found: list[str] = []
    for p in root.rglob("*"):
        if p.is_file() and p.name == target:
            found.append(str(p.relative_to(root)).replace("\\", "/"))
    return found


DELIVERABLE_EXTS = {
    ".xlsx", ".xls", ".xlsm", ".csv", ".sql", ".pdf", ".zip",
    ".md", ".json", ".docx", ".pptx", ".html", ".htm",
}


def list_deliverable_files(sandbox_id: str, limit: int = 20) -> list[str]:
    """Final deliverable files at the workplace root (exclude task/ & tmp/ scratch)."""
    root = ensure_workplace(sandbox_id)
    if not root.is_dir():
        return []
    files: list[tuple[float, str]] = []
    for p in root.iterdir():
        if not p.is_file() or p.name.startswith("."):
            continue
        if p.suffix.lower() not in DELIVERABLE_EXTS:
            continue
        try:
            mtime = p.stat().st_mtime
        except OSError:
            continue
        files.append((mtime, p.name))
    files.sort(key=lambda x: x[0], reverse=True)
    return [name for _, name in files[: max(1, limit)]]


def _cell_value(value) -> str:
    if value is None:
        return ""
    return str(value)


def _normalize_sheet_rows(rows: list[list]) -> list[list[str]]:
    if not rows:
        return [[""]]
    max_cols = max(len(row) for row in rows)
    return [
        [_cell_value(cell) for cell in row] + [""] * (max_cols - len(row))
        for row in rows
    ]


def _read_excel_preview(target: Path) -> dict:
    import openpyxl

    wb = openpyxl.load_workbook(target, read_only=True, data_only=True)
    sheets = []
    truncated = False
    try:
        for name in wb.sheetnames:
            ws = wb[name]
            rows = []
            for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
                if row_idx >= MAX_EXCEL_ROWS:
                    truncated = True
                    break
                cells = [_cell_value(cell) for cell in row[:MAX_EXCEL_COLS]]
                if len(row) > MAX_EXCEL_COLS:
                    truncated = True
                rows.append(cells)
            sheets.append({"name": name, "rows": _normalize_sheet_rows(rows)})
    finally:
        wb.close()
    return {"ok": True, "type": "excel", "sheets": sheets, "truncated": truncated}


def _save_excel_workbook(target: Path, sheets_json: str) -> dict:
    import openpyxl

    try:
        sheets = json.loads(sheets_json or "[]")
    except json.JSONDecodeError:
        return {"ok": False, "msg": "无效的表格数据"}
    if not isinstance(sheets, list):
        return {"ok": False, "msg": "无效的表格数据"}

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet in sheets:
        if not isinstance(sheet, dict):
            continue
        title = str(sheet.get("name") or "Sheet1")[:31] or "Sheet1"
        ws = wb.create_sheet(title=title)
        for row_idx, row in enumerate(sheet.get("rows") or [], 1):
            if not isinstance(row, list):
                continue
            for col_idx, value in enumerate(row[:MAX_EXCEL_COLS], 1):
                ws.cell(row=row_idx, column=col_idx, value="" if value is None else str(value))
    if not wb.sheetnames:
        wb.create_sheet("Sheet1")
    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)
    return {"ok": True}


def wp_action(sandbox_id: str, action: str, body) -> dict:
    root = _resolved_wp_root(sandbox_id)
    path = body.path or ""
    target = _path_under_root(root, path.lstrip("/"))
    if target is None:
        return {"ok": False, "msg": "非法路径"}

    if action == "mkdir_workplace":
        target.mkdir(parents=True, exist_ok=True)
        return {"ok": True}
    if action == "delete_workplace":
        if target.is_dir():
            shutil.rmtree(target)
        else:
            target.unlink(missing_ok=True)
        return {"ok": True}
    if action == "rename_workplace":
        new_name = body.new_name or body.name or ""
        if "/" not in new_name.lstrip("/"):
            if target == root:
                new_rel = new_name
            else:
                new_rel = f"{target.parent.relative_to(root)}/{new_name}".replace("\\", "/")
        else:
            new_rel = new_name.lstrip("/")
        new_path = _path_under_root(root, new_rel)
        if new_path is None:
            return {"ok": False, "msg": "非法路径"}
        target.rename(new_path)
        return {"ok": True}
    if action == "move_workplace":
        dest = _path_under_root(root, (body.dest or "").lstrip("/"))
        if dest is None:
            return {"ok": False, "msg": "非法路径"}
        dest.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(target), str(dest))
        return {"ok": True}
    if action == "view_workplace":
        if not target.is_file():
            return {"ok": False, "msg": "不是文件"}
        ext = target.suffix.lower()
        if ext in EXCEL_EXTS:
            try:
                return _read_excel_preview(target)
            except Exception as exc:
                return {"ok": False, "msg": f"无法读取 Excel：{exc}"}
        if ext == ".xls":
            return {
                "ok": False,
                "type": "excel_legacy",
                "msg": "旧版 .xls 文件请下载后使用 Excel 打开",
            }
        return {
            "ok": True,
            "type": "text",
            "content": target.read_text(encoding="utf-8", errors="replace")[:50000],
        }
    if action == "save_workplace":
        if target.suffix.lower() in EXCEL_EXTS:
            return {"ok": False, "msg": "Excel 文件请使用表格编辑器保存"}
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(body.content or "", encoding="utf-8")
        return {"ok": True}
    if action == "save_excel_workplace":
        if target.suffix.lower() not in EXCEL_EXTS:
            return {"ok": False, "msg": "不是 Excel 文件"}
        return _save_excel_workbook(target, body.content or "[]")
    if action == "zip_workplace":
        zip_path = target.with_suffix(".zip") if target.is_file() else root / f"{target.name}.zip"
        with zipfile.ZipFile(zip_path, "w") as zf:
            if target.is_dir():
                for f in target.rglob("*"):
                    if f.is_file():
                        zf.write(f, f.relative_to(target.parent))
            else:
                zf.write(target, target.name)
        return {"ok": True, "zip": str(zip_path.name)}
    if action == "unzip_workplace":
        with zipfile.ZipFile(target, "r") as zf:
            zf.extractall(target.parent)
        return {"ok": True}
    return {"ok": False, "msg": "未知操作"}


def download_path(sandbox_id: str, rel: str) -> Path | None:
    root = _resolved_wp_root(sandbox_id)
    clean = str(rel or "").strip().lstrip("/")
    if clean.startswith("workplace/"):
        clean = clean[len("workplace/") :]
    p = _path_under_root(root, clean)
    if p and p.is_file():
        return p
    name = clean.rsplit("/", 1)[-1] if "/" in clean else clean
    if name:
        matches = find_files(sandbox_id, name)
        if matches:
            def _mtime(rel_path: str) -> float:
                fp = _path_under_root(root, rel_path)
                try:
                    return fp.stat().st_mtime if fp else 0.0
                except OSError:
                    return 0.0
            matches.sort(key=_mtime, reverse=True)
            return _path_under_root(root, matches[0])
    return None


def upload_file(sandbox_id: str, rel_dir: str, filename: str, content: bytes) -> dict:
    root = _resolved_wp_root(sandbox_id)
    root.mkdir(parents=True, exist_ok=True)
    target_dir = _path_under_root(root, rel_dir.lstrip("/"))
    if target_dir is None:
        return {"ok": False, "msg": "非法路径"}
    target_dir.mkdir(parents=True, exist_ok=True)
    safe_name = _safe_filename(filename)
    if not safe_name:
        return {"ok": False, "msg": "无效文件名"}
    target = _path_under_root(target_dir, safe_name)
    if target is None:
        return {"ok": False, "msg": "非法路径"}
    target.write_bytes(content)
    rel_path = str(target.relative_to(root)).replace("\\", "/")
    return {"ok": True, "path": rel_path}
