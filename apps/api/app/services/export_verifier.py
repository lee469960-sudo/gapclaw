"""Pre-FINAL export verifier: file / sheets / columns / core nodes / summary."""

from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path
from typing import Any

from app.services.workplace import ensure_workplace

_COHORT_ROW_COMPLETE_RATIO = 0.85


_QUERY_CHANGE_DIAGNOSIS_SYSTEM = """你是数据查询差异审计器。只输出一个 JSON 对象，不要 Markdown 或思考过程。

根据用户质疑、上一版 SQL、当前 SQL、本轮 MCP schema 备注和引擎提供的真实运行证据，解释结果为何变化：
{"summary":"一句具体结论；证据不足时明确说明缺什么", "changes":[{"subject":"变化点", "before":"上一版真实条件/表达式", "after":"当前版真实条件/表达式", "cause":"该变化如何影响结果；区分已验证事实和基于 SQL/schema 的推断"}]}

规则：
1. 不得编造人数、金额、字段含义或执行结果；数字只能引用 runtime_evidence。
2. before/after 必须来自所给 SQL；没有上一版 SQL 时明确写“无可比较基线”。
3. schema 备注可以解释字段语义，但应标明是备注证据；无法确定根因时说明需要哪条对照查询。
4. changes 只列与 verification_targets 和用户反馈相关的实质差异，不复述整条 SQL。
"""


def _first_json_object(text: str) -> dict[str, Any] | None:
    raw = str(text or "")
    decoder = json.JSONDecoder()
    for index, char in enumerate(raw):
        if char != "{":
            continue
        try:
            value, _ = decoder.raw_decode(raw[index:])
        except (TypeError, ValueError, json.JSONDecodeError):
            continue
        if isinstance(value, dict):
            return value
    return None


async def diagnose_query_export_change(
    llm,
    *,
    prior_sql: str,
    current_sql: str,
    schemas: dict[str, Any] | None,
    verification: dict[str, Any] | None,
    verification_targets: list[str] | None,
    user_context: str = "",
    db=None,
    timeout: int = 45,
) -> dict[str, Any]:
    """Ask the model for a causal SQL diff while keeping runtime facts authoritative."""
    from app.services.llm_client import chat_completion

    facts = verification if isinstance(verification, dict) else {}
    comparison = facts.get("comparison") if isinstance(facts.get("comparison"), dict) else {}
    runtime_evidence = {
        "expected_row_count": facts.get("expected_row_count"),
        "fetched_row_count": facts.get("fetched_row_count"),
        "file_row_count": facts.get("file_row_count"),
        "row_count_before": comparison.get("row_count_before"),
        "row_count_after": comparison.get("row_count_after"),
        "data_changed": comparison.get("data_changed"),
    }
    payload = {
        "verification_targets": [
            str(x).strip() for x in (verification_targets or []) if str(x).strip()
        ],
        "user_context": str(user_context or "")[:4000],
        "prior_sql": str(prior_sql or "")[:20000],
        "current_sql": str(current_sql or "")[:20000],
        "schemas": schemas or {},
        "runtime_evidence": runtime_evidence,
    }
    raw = await chat_completion(
        llm,
        [
            {"role": "system", "content": _QUERY_CHANGE_DIAGNOSIS_SYSTEM},
            {"role": "user", "content": json.dumps(payload, ensure_ascii=False)},
        ],
        max_tokens=1800,
        db=db,
        timeout=timeout,
    )
    data = _first_json_object(raw) or {}
    summary = str(data.get("summary") or "").strip()
    changes: list[dict[str, str]] = []
    for row in data.get("changes") or []:
        if not isinstance(row, dict):
            continue
        item = {
            key: str(row.get(key) or "").strip()
            for key in ("subject", "before", "after", "cause")
        }
        if item["subject"] and item["cause"]:
            changes.append(item)
    if not summary:
        summary = "模型未返回可用的差异结论；需要补充针对质疑点的对照查询。"
    return {
        "summary": summary,
        "changes": changes[:8],
        "runtime_evidence": runtime_evidence,
    }


def inspect_query_deliverable(path: Path | None) -> dict[str, Any]:
    """Read generic delivery facts and a stable data-sheet fingerprint."""
    if path is None or not path.is_file():
        return {"exists": False, "headers": [], "row_count": None, "fingerprint": ""}
    try:
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as ex:
        return {
            "exists": True,
            "headers": [],
            "row_count": None,
            "fingerprint": "",
            "error": f"{type(ex).__name__}: {ex}",
        }
    try:
        ws = next(
            (wb[name] for name in wb.sheetnames if name == "数据" or "数据" in name),
            None,
        )
        if ws is None:
            ws = next(
                (wb[name] for name in wb.sheetnames if "口径" not in name),
                None,
            )
        if ws is None:
            return {"exists": True, "headers": [], "row_count": 0, "fingerprint": ""}
        rows = ws.iter_rows(values_only=True)
        try:
            first = next(rows)
        except StopIteration:
            return {"exists": True, "headers": [], "row_count": 0, "fingerprint": ""}
        headers = [str(value or "").strip() for value in first]
        digest = hashlib.sha256()
        row_count = 0
        for row in rows:
            if not row or all(value is None or str(value).strip() == "" for value in row):
                continue
            row_count += 1
            digest.update(
                json.dumps(
                    list(row), ensure_ascii=False, default=str, separators=(",", ":"),
                ).encode("utf-8")
            )
            digest.update(b"\n")
        return {
            "exists": True,
            "headers": headers,
            "row_count": row_count,
            "fingerprint": digest.hexdigest(),
        }
    finally:
        wb.close()


def verify_query_export_delivery(
    *,
    current: dict[str, Any],
    requested_columns: list[str],
    expected_row_count: int | None,
    fetched_row_count: int,
    pages_fetched: int,
    mcp_call_count: int,
    prior: dict[str, Any] | None = None,
    verification_targets: list[str] | None = None,
) -> dict[str, Any]:
    """Verify query execution and xlsx materialization using runtime facts only."""
    headers = [str(x) for x in (current.get("headers") or [])]
    file_rows = current.get("row_count")
    checks = {
        "file_exists": bool(current.get("exists")),
        "columns_match": headers == list(requested_columns),
        "count_matches_fetch": expected_row_count == fetched_row_count,
        "fetch_matches_file": fetched_row_count == file_rows,
        "has_fingerprint": bool(current.get("fingerprint")) or fetched_row_count == 0,
    }
    previous = prior if isinstance(prior, dict) and prior.get("exists") else None
    comparison = {
        "available": bool(previous),
        "row_count_before": previous.get("row_count") if previous else None,
        "row_count_after": file_rows,
        "headers_changed": (
            previous.get("headers") != headers if previous else None
        ),
        "data_changed": (
            previous.get("fingerprint") != current.get("fingerprint")
            if previous else None
        ),
    }
    return {
        "status": "pass" if all(checks.values()) else "failed",
        "checks": checks,
        "requested_columns": list(requested_columns),
        "expected_row_count": expected_row_count,
        "fetched_row_count": fetched_row_count,
        "file_row_count": file_rows,
        "pages_fetched": pages_fetched,
        "mcp_call_count": mcp_call_count,
        "fingerprint": str(current.get("fingerprint") or ""),
        "comparison": comparison,
        "verification_targets": [
            str(x).strip() for x in (verification_targets or []) if str(x).strip()
        ],
    }


def format_query_export_verification(result: dict[str, Any] | None) -> str:
    data = result if isinstance(result, dict) else {}
    if not data:
        return ""
    ok = data.get("status") == "pass"
    lines = [
        "### 严格验证",
        "",
        f"- 结果：{'通过' if ok else '未通过'}",
        (
            f"- MCP：共调用 {int(data.get('mcp_call_count') or 0)} 次，"
            f"COUNT={data.get('expected_row_count')}，分页 {int(data.get('pages_fetched') or 0)} 页"
        ),
        (
            f"- 行数：查询取得 {data.get('fetched_row_count')} 行，"
            f"xlsx 数据 sheet {data.get('file_row_count')} 行"
        ),
        (
            "- 列顺序："
            + ("与用户请求完全一致" if (data.get("checks") or {}).get("columns_match") else "不一致")
        ),
    ]
    targets = [str(x) for x in (data.get("verification_targets") or []) if str(x)]
    if targets:
        lines.append("- 本轮核验重点：" + "；".join(targets))
    comparison = data.get("comparison") if isinstance(data.get("comparison"), dict) else {}
    if comparison.get("available"):
        lines.append(
            f"- 与上一份交付比较：行数 {comparison.get('row_count_before')} → "
            f"{comparison.get('row_count_after')}，数据内容"
            + ("有变化" if comparison.get("data_changed") else "一致")
        )
    diagnosis = data.get("diagnosis") if isinstance(data.get("diagnosis"), dict) else {}
    if diagnosis:
        lines.extend(["", "### 差异原因", "", str(diagnosis.get("summary") or "").strip()])
        for row in diagnosis.get("changes") or []:
            if not isinstance(row, dict):
                continue
            subject = str(row.get("subject") or "变化点").strip()
            before = str(row.get("before") or "无可比较基线").strip()
            after = str(row.get("after") or "未说明").strip()
            cause = str(row.get("cause") or "").strip()
            line = f"- {subject}：`{before}` → `{after}`"
            if cause:
                line += f"；{cause}"
            lines.append(line)
    return "\n".join(lines)


def _sheet_names(path: Path) -> list[str]:
    try:
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        names = list(wb.sheetnames)
        wb.close()
        return names
    except Exception:
        return []


def _read_headers_and_rows(
    path: Path,
) -> tuple[list[str], int, dict[str, Any], dict[str, Any], dict[str, int]]:
    """Return headers, data row count, money sums, numeric sums, nonzero counts."""
    try:
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return [], 0, {}, {}, {}
    try:
        # Prefer「数据」then first non-口径 sheet
        ws = None
        for name in wb.sheetnames:
            if name in ("数据", "Sheet1", "充值用户数据") or "数据" in name:
                ws = wb[name]
                break
        if ws is None:
            for name in wb.sheetnames:
                if "口径" in name:
                    continue
                ws = wb[name]
                break
        if ws is None:
            return [], 0, {}, {}, {}
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            return [], 0, {}, {}, {}
        headers = [str(c or "").strip() for c in header_row]
        n = 0
        money_sums: dict[str, float] = {}
        numeric_sums: dict[str, float] = {}
        nonzero_counts: dict[str, int] = {}
        money_idx = [
            i
            for i, h in enumerate(headers)
            if any(k in h for k in ("充值金额", "提现金额", "下注金额", "返奖金额"))
        ]
        for row in rows_iter:
            if not row or all(v is None or str(v).strip() == "" for v in row):
                continue
            n += 1
            for i in money_idx:
                if i >= len(row):
                    continue
                try:
                    money_sums[headers[i]] = money_sums.get(headers[i], 0.0) + float(row[i] or 0)
                except (TypeError, ValueError):
                    pass
            for i, h in enumerate(headers):
                if i >= len(row):
                    continue
                try:
                    val = float(row[i] or 0)
                except (TypeError, ValueError):
                    continue
                numeric_sums[h] = numeric_sums.get(h, 0.0) + val
                if abs(val) > 1e-12:
                    nonzero_counts[h] = nonzero_counts.get(h, 0) + 1
        return headers, n, money_sums, numeric_sums, nonzero_counts
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _has_koujing_sheet(names: list[str]) -> bool:
    return any("口径" in n for n in names)


def _has_data_sheet(names: list[str]) -> bool:
    if any(n in ("数据", "Sheet1") or "数据" in n for n in names):
        return True
    # single-sheet deliverable still OK if not only 口径
    non_kj = [n for n in names if "口径" not in n]
    return bool(non_kj)


def _verifier_status(
    *,
    errors: list[str],
    warnings: list[str],
    missing_core: list[str],
    missing_columns: list[str],
    incomplete_non_core: list[str],
) -> str:
    if errors or missing_core:
        return "failed"
    if warnings or missing_columns or incomplete_non_core:
        return "repairable"
    return "pass"


def _repair_hints(
    *,
    missing_core: list[str],
    missing_columns: list[str],
    incomplete_non_core: list[str],
    errors: list[str],
) -> list[str]:
    hints: list[str] = []
    if missing_core:
        hints.append("先补拉核心 role：" + "、".join(missing_core))
    if missing_columns:
        hints.append("重跑写表或修正列计划，补齐缺失列：" + "、".join(missing_columns[:6]))
    if incomplete_non_core:
        hints.append("补拉非核心缺口，或在口径说明中保留未完整标记：" + "、".join(incomplete_non_core[:6]))
    if any("交付文件" in e or "数据行为 0" in e or "数据 sheet" in e for e in errors):
        hints.append("重新物化交付文件，并确认数据 sheet 存在且非空。")
    return list(dict.fromkeys(hints))


def _schema_discovery_from_trace(trace: dict[str, Any]) -> dict[str, Any]:
    schema = trace.get("schema_discovery")
    if isinstance(schema, dict):
        return schema
    contract = trace.get("export_contract")
    if isinstance(contract, dict) and isinstance(contract.get("schema_discovery"), dict):
        return dict(contract["schema_discovery"])
    return {}


def _planned_schema_views(plan: list[dict[str, Any]]) -> list[str]:
    views: list[str] = []
    for c in plan:
        if not isinstance(c, dict):
            continue
        spec = c.get("agg_spec") if isinstance(c.get("agg_spec"), dict) else {}
        view = str(spec.get("view") or "").strip()
        if not view:
            src = str(c.get("数据来源") or "")
            m = re.search(r"\bview_[A-Za-z0-9_]+\b", src)
            view = m.group(0) if m else ""
        if view:
            views.append(view)
    return list(dict.fromkeys(views))


def _append_schema_discovery_warnings(
    *,
    warnings: list[str],
    summary: dict[str, Any],
    plan: list[dict[str, Any]],
    trace: dict[str, Any],
) -> None:
    schema = _schema_discovery_from_trace(trace)
    if not schema or not schema.get("required"):
        return
    listed = schema.get("list_ads_views")
    listed_ok = bool(isinstance(listed, dict) and listed.get("captured"))
    described = schema.get("described_views")
    described_views = set((described or {}).keys()) if isinstance(described, dict) else set()
    planned_views = _planned_schema_views(plan)
    missing_describe = [v for v in planned_views if v not in described_views]
    summary["schema_discovery"] = {
        "required": True,
        "list_ads_views": listed_ok,
        "described_views": sorted(described_views),
        "missing_describe_views": missing_describe,
    }
    if not listed_ok:
        warnings.append("缺少 MCP 接口列表发现证据: list_ads_views")
    if missing_describe:
        warnings.append(
            "缺少 MCP 表备注/字段发现证据: describe_ads_view "
            + "、".join(missing_describe[:6])
        )


def _cohort_uid_estimate_from(task_spec: dict[str, Any], trace: dict[str, Any]) -> int | None:
    candidates = [
        task_spec.get("cohort_uid_estimate") if isinstance(task_spec, dict) else None,
        trace.get("cohort_uid_estimate") if isinstance(trace, dict) else None,
    ]
    if isinstance(trace, dict):
        contract = trace.get("export_contract")
        if isinstance(contract, dict):
            candidates.append(contract.get("cohort_uid_estimate"))
        trace_spec = trace.get("task_spec")
        if isinstance(trace_spec, dict):
            candidates.append(trace_spec.get("cohort_uid_estimate"))
    for raw in candidates:
        try:
            n = int(raw)
        except (TypeError, ValueError):
            continue
        if n > 0:
            return n
    return None


def _finalize_result(
    *,
    errors: list[str],
    warnings: list[str],
    summary: dict[str, Any],
    missing_core: list[str] | None = None,
    missing_columns: list[str] | None = None,
    incomplete_non_core: list[str] | None = None,
) -> dict[str, Any]:
    miss_core = list(dict.fromkeys(missing_core or []))
    miss_cols = list(dict.fromkeys(missing_columns or []))
    incomplete = list(dict.fromkeys(incomplete_non_core or []))
    status = _verifier_status(
        errors=errors,
        warnings=warnings,
        missing_core=miss_core,
        missing_columns=miss_cols,
        incomplete_non_core=incomplete,
    )
    blocking = list(dict.fromkeys([*errors, *(f"核心缺失: {r}" for r in miss_core)]))
    hints = _repair_hints(
        missing_core=miss_core,
        missing_columns=miss_cols,
        incomplete_non_core=incomplete,
        errors=errors,
    )
    summary["status"] = status
    summary["passed"] = status == "pass"
    summary["missing_core"] = miss_core
    summary["missing_columns"] = miss_cols
    summary["incomplete_non_core"] = incomplete
    return {
        # Legacy: no hard error. New callers should prefer status / ok.
        "passed": status != "failed",
        "ok": status == "pass",
        "status": status,
        "errors": errors,
        "warnings": warnings,
        "missing_core": miss_core,
        "missing_columns": miss_cols,
        "incomplete_non_core": incomplete,
        "summary": summary,
        "repair_hints": hints,
        "blocking_reasons": blocking,
    }


def verify_export_deliverable(
    *,
    sandbox_id: str,
    file_rel: str,
    task_spec: dict | None = None,
    column_plan: list[dict] | None = None,
    trace: dict | None = None,
    failed_views: set[str] | list[str] | None = None,
    abandoned_roles: set[str] | list[str] | None = None,
) -> dict[str, Any]:
    """Return standardized VerifierResult.

    Legacy `passed` is kept for compatibility. New completion checks should use
    `status == "pass"` or `ok`.
    """
    errors: list[str] = []
    warnings: list[str] = []
    missing_core: list[str] = []
    missing_columns: list[str] = []
    incomplete_non_core: list[str] = []
    summary: dict[str, Any] = {
        "file_rel": file_rel or "",
        "file_size": None,
        "row_count": 0,
        "headers": [],
        "sheet_names": [],
        "money_sums": {},
        "numeric_sums": {},
        "nonzero_counts": {},
        "incomplete_columns": [],
        "schema_discovery": {},
    }
    spec = task_spec if isinstance(task_spec, dict) else {}
    plan = [c for c in (column_plan or []) if isinstance(c, dict)]
    tr = trace if isinstance(trace, dict) else {}
    failed = {str(x) for x in (failed_views or tr.get("failed_views") or [])}
    abandoned = {str(x) for x in (abandoned_roles or tr.get("abandoned_roles") or [])}

    for node in tr.get("query_nodes") or []:
        if not isinstance(node, dict):
            continue
        role = str(node.get("role") or "").strip().lower()
        st = str(node.get("status") or "").strip().lower()
        if role in ("user", "pay") and st in ("error", "soft_fail") and not node.get("fallback_used"):
            missing_core.append(role)
    if "user" in abandoned or "pay" in abandoned:
        missing_core.extend(sorted(abandoned & {"user", "pay"}))

    if not file_rel or not sandbox_id:
        errors.append("缺少交付路径或 sandbox")
        return _finalize_result(
            errors=errors,
            warnings=warnings,
            summary=summary,
            missing_core=missing_core,
            missing_columns=missing_columns,
            incomplete_non_core=incomplete_non_core,
        )

    root = ensure_workplace(sandbox_id)
    path = root / str(file_rel).lstrip("/")
    if not path.is_file():
        errors.append(f"交付文件不存在: {file_rel}")
        return _finalize_result(
            errors=errors,
            warnings=warnings,
            summary=summary,
            missing_core=missing_core,
            missing_columns=missing_columns,
            incomplete_non_core=incomplete_non_core,
        )

    size = path.stat().st_size
    summary["file_size"] = size
    if size <= 0:
        errors.append("交付文件为空")

    names = _sheet_names(path)
    summary["sheet_names"] = names
    if not _has_data_sheet(names):
        errors.append("缺少数据 sheet")
    if not _has_koujing_sheet(names):
        warnings.append("缺少口径说明 sheet（兼容单 sheet 时降级为警告）")

    headers, row_count, money_sums, numeric_sums, nonzero_counts = _read_headers_and_rows(path)
    summary["headers"] = headers
    summary["row_count"] = row_count
    summary["money_sums"] = {k: round(v, 2) for k, v in money_sums.items()}
    summary["numeric_sums"] = {k: round(v, 2) for k, v in numeric_sums.items()}
    summary["nonzero_counts"] = dict(nonzero_counts)
    cohort_uid_estimate = _cohort_uid_estimate_from(spec, tr)
    if cohort_uid_estimate:
        summary["cohort_uid_estimate"] = cohort_uid_estimate
    if row_count <= 0:
        errors.append("数据行为 0")
    elif cohort_uid_estimate and row_count < int(cohort_uid_estimate * _COHORT_ROW_COMPLETE_RATIO):
        warnings.append(
            f"交付行数 {row_count} 低于 cohort 目标≈{cohort_uid_estimate}"
        )
        incomplete_non_core.append("数据行数")

    expected = [
        str(c.get("header") or "").strip()
        for c in plan
        if str(c.get("header") or "").strip()
    ]
    if not expected:
        expected = [str(x).strip() for x in (spec.get("requested_columns") or []) if str(x).strip()]
    if expected:
        # Order: prefix match — allow parenthetical suffixes on file headers
        for i, exp in enumerate(expected):
            if i >= len(headers):
                errors.append(f"缺少输出列: {exp}")
                missing_columns.append(exp)
                continue
            stem = re.split(r"[（(]", exp, maxsplit=1)[0].strip()
            hstem = re.split(r"[（(]", headers[i], maxsplit=1)[0].strip()
            if stem and stem not in headers[i] and hstem not in exp and stem != hstem:
                warnings.append(f"列顺序/名称可能漂移: 期望「{exp}」实际「{headers[i]}」")

    # Core query nodes must not claim success if failed
    for node in tr.get("query_nodes") or []:
        if not isinstance(node, dict):
            continue
        role = str(node.get("role") or "").strip().lower()
        st = str(node.get("status") or "").strip().lower()
        policy = "optional"
        for c in plan:
            if str((c.get("agg_spec") or {}).get("role") or "").lower() == role:
                policy = str(c.get("failure_policy") or policy)
        if role in ("user", "pay") or policy == "core":
            if st in ("error", "soft_fail") and not node.get("fallback_used"):
                # soft_fail on core without fallback is hard error
                if role in ("user", "pay"):
                    errors.append(
                        f"核心节点失败不可伪造成功: {node.get('key') or role}/{node.get('view')}"
                    )

    if "user" in abandoned or "pay" in abandoned:
        core_abandoned = sorted(abandoned & {"user", "pay"})
        errors.append(f"核心角色已放弃: {core_abandoned}")

    # Optional/heavy incomplete → must be annotated in plan 统计方法
    incomplete_cols: list[str] = []
    for c in plan:
        header = str(c.get("header") or "")
        method = str(c.get("统计方法") or "")
        policy = str(c.get("failure_policy") or "")
        role = str((c.get("agg_spec") or {}).get("role") or "").strip().lower()
        mode = str(c.get("fetch_mode") or "").lower()
        view = str((c.get("agg_spec") or {}).get("view") or "")
        failed_hit = (
            view in failed
            or f"node:" in "".join(failed)
            or role in abandoned
        )
        if policy in ("soft_fail", "optional") and failed_hit:
            if "未完整" not in method and "未完整" not in str(c.get("口径") or ""):
                warnings.append(f"失败列未标未完整: {header}")
            incomplete_cols.append(header)
            incomplete_non_core.append(header)
        if mode in ("sequence", "top_n") and failed_hit and "未完整" not in method:
            warnings.append(f"heavy 列失败建议标未完整: {header}")
            incomplete_non_core.append(header)
    summary["incomplete_columns"] = incomplete_cols

    # Unexplained empty money columns
    for h in headers:
        if "金额" in h and h in money_sums and money_sums[h] == 0:
            if h not in incomplete_cols and "未完整" not in "".join(
                str(c.get("统计方法") or "") for c in plan if h in str(c.get("header") or "")
            ):
                warnings.append(f"金额列全 0 且未标缺口: {h}")
    for h in headers:
        if not any(k in h for k in ("连续充值次数", "下注次数", "流水倍数")):
            continue
        if h in numeric_sums and numeric_sums[h] == 0 and row_count > 0:
            if h not in incomplete_cols:
                warnings.append(f"指标列全 0 且未标缺口: {h}")
                incomplete_cols.append(h)
                incomplete_non_core.append(h)
    summary["incomplete_columns"] = list(dict.fromkeys(incomplete_cols))

    _append_schema_discovery_warnings(
        warnings=warnings,
        summary=summary,
        plan=plan,
        trace=tr,
    )

    return _finalize_result(
        errors=errors,
        warnings=warnings,
        summary=summary,
        missing_core=missing_core,
        missing_columns=missing_columns,
        incomplete_non_core=incomplete_non_core,
    )


def format_verifier_block(result: dict | None) -> str:
    """Markdown block for FINAL appendix / failure message."""
    r = result if isinstance(result, dict) else {}
    sm = r.get("summary") if isinstance(r.get("summary"), dict) else {}
    lines = ["### 校验摘要"]
    status = str(r.get("status") or ("pass" if r.get("passed") else "failed"))
    lines.append(f"- 状态: {status}")
    lines.append(f"- 通过: {'是' if r.get('ok') else '否'}")
    if sm.get("row_count") is not None:
        lines.append(f"- 数据行数: {sm.get('row_count')}")
    if sm.get("file_size") is not None:
        lines.append(f"- 文件大小: {sm.get('file_size')} 字节")
    schema = sm.get("schema_discovery") if isinstance(sm.get("schema_discovery"), dict) else {}
    if schema.get("required"):
        lines.append(
            "- Schema 发现: "
            + ("list_ads_views 已记录" if schema.get("list_ads_views") else "缺 list_ads_views")
        )
        missing_views = [
            str(v)
            for v in (schema.get("missing_describe_views") or [])
            if str(v).strip()
        ]
        if missing_views:
            lines.append(
                "- Schema 缺口: 缺 describe_ads_view "
                + "、".join(missing_views[:6])
            )
    for e in (r.get("errors") or [])[:8]:
        lines.append(f"- 错误: {e}")
    for w in (r.get("warnings") or [])[:8]:
        lines.append(f"- 警告: {w}")
    for h in (r.get("repair_hints") or [])[:5]:
        lines.append(f"- 修复建议: {h}")
    return "\n".join(lines)


def build_final_summary_from_verifier(
    result: dict | None,
    *,
    time_label: str = "",
) -> str:
    """Deterministic 导出概况 markdown from verifier.summary (no LLM numbers)."""
    r = result if isinstance(result, dict) else {}
    sm = r.get("summary") if isinstance(r.get("summary"), dict) else {}
    lines = [
        "### 导出概况",
        "| 指标 | 数值 |",
        "| --- | --- |",
        f"| 校验状态 | {str(r.get('status') or sm.get('status') or 'unknown')} |",
        f"| 时间范围 | {time_label or '见需求'} |",
        f"| 总用户数 | {int(sm.get('row_count') or 0):,} |",
    ]
    money = sm.get("money_sums") if isinstance(sm.get("money_sums"), dict) else {}
    for k, v in list(money.items())[:6]:
        lines.append(f"| {k} | {v} |")
    if sm.get("file_size") is not None:
        size = int(sm["file_size"])
        if size >= 1024 * 1024:
            human = f"{size / (1024 * 1024):.1f} MB"
        elif size >= 1024:
            human = f"{size / 1024:.1f} KB"
        else:
            human = f"{size} B"
        lines.append(f"| 文件大小 | {human} |")
    if sm.get("incomplete_columns"):
        lines.append(
            f"| 未完整列 | {'、'.join(sm['incomplete_columns'][:8])} |"
        )
    return "\n".join(lines)
