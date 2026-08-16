"""Platform merge + dual-sheet export deliverable (Codex-style, column-plan driven).

Generic: any column_plan + task pages → sheet「数据」+ sheet「口径说明」.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from app.services.export_column_plan import column_source_method
from app.services.workplace import (
    ensure_workplace,
    materialize_analyzed_export,
)


def _safe_stem(name: str) -> str:
    s = re.sub(r"[^\w\u4e00-\u9fff\-]+", "_", (name or "").strip())
    return (s[:80] or "export").strip("_")


def export_filename_for_title(title: str, *, run_id: str = "") -> str:
    """Build a unique workbook name from the model-understood task title."""
    stem = _safe_stem(title)
    rid = re.sub(r"[^A-Za-z0-9_-]+", "", str(run_id or ""))[:48]
    return f"{stem}_{rid}.xlsx" if rid else f"{stem}.xlsx"


def _append_koujing_sheet(
    dest: Path,
    column_plan: list[dict] | None,
    *,
    title: str = "",
) -> None:
    """Add/replace「口径说明」sheet on an existing xlsx."""
    if not dest.is_file():
        return
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    except ImportError:
        return
    try:
        wb = openpyxl.load_workbook(dest)
    except Exception:
        return
    sheet_name = "口径说明"
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    headers = ["#", "列名", "数据来源", "统计方法"]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    thin = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = thin
    plan = [c for c in (column_plan or []) if isinstance(c, dict) and c.get("header")]
    for i, col in enumerate(plan, 1):
        name = str(col.get("header") or "")
        src, method = column_source_method(col)
        for c, val in enumerate([i, name, src, method], 1):
            cell = ws.cell(i + 1, c, val)
            cell.border = thin
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    if title:
        ws.cell(len(plan) + 3, 1, f"任务：{title}")
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 42
    ws.column_dimensions["D"].width = 48
    # Prefer data sheet first
    if wb.sheetnames and wb.sheetnames[0] == sheet_name and len(wb.sheetnames) > 1:
        wb.move_sheet(wb[wb.sheetnames[1]], offset=-1)
    try:
        wb.save(dest)
    except Exception:
        pass
    finally:
        wb.close()


def _style_data_sheet(dest: Path) -> None:
    """Freeze panes + filter on first sheet (Codex Excel polish)."""
    if not dest.is_file():
        return
    try:
        import openpyxl
    except ImportError:
        return
    try:
        wb = openpyxl.load_workbook(dest)
        ws = wb.active
        if ws.max_row >= 1 and ws.max_column >= 1:
            ws.auto_filter.ref = ws.dimensions
            ws.freeze_panes = "C2" if ws.max_column >= 2 else "A2"
        wb.save(dest)
        wb.close()
    except Exception:
        pass


def write_export_deliverable(
    sandbox_id: str,
    run_id: str,
    *,
    column_plan: list[dict] | None = None,
    time_window: dict | None = None,
    column_headers: list[str] | None = None,
    preferred_name: str = "",
    title: str = "",
) -> str | None:
    """Merge task pages into dual-sheet xlsx; return workplace-relative path.

    Uses existing join materializer for row data, then adds 口径说明 sheet.
    """
    headers = column_headers
    if not headers and column_plan:
        headers = [
            str(c.get("header") or "").strip()
            for c in column_plan
            if isinstance(c, dict) and str(c.get("header") or "").strip()
        ]
    name = preferred_name or ""
    if not name and title:
        name = f"{_safe_stem(title)}.xlsx"
    rel = materialize_analyzed_export(
        sandbox_id,
        run_id,
        time_window=time_window,
        column_headers=headers,
        preferred_name=name,
        column_plan=column_plan,
    )
    if not rel:
        return None
    root = ensure_workplace(sandbox_id)
    dest = root / rel
    if not dest.is_file():
        return rel
    try:
        import openpyxl
        wb = openpyxl.load_workbook(dest)
        active = wb.active
        if active is not None and active.title in ("Sheet", "Sheet1", "用户分析"):
            active.title = "数据"
        # Ensure 口径说明
        sheet_name = "口径说明"
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]
        ws = wb.create_sheet(sheet_name)
        from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="305496")
        thin = Border(
            left=Side(style="thin", color="BFBFBF"),
            right=Side(style="thin", color="BFBFBF"),
            top=Side(style="thin", color="BFBFBF"),
            bottom=Side(style="thin", color="BFBFBF"),
        )
        for c, h in enumerate(["#", "列名", "数据来源", "统计方法"], 1):
            cell = ws.cell(1, c, h)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin
        plan = [c for c in (column_plan or []) if isinstance(c, dict) and c.get("header")]
        for i, col in enumerate(plan, 1):
            src, method = column_source_method(col)
            for ci, val in enumerate([i, str(col.get("header") or ""), src, method], 1):
                cell = ws.cell(i + 1, ci, val)
                cell.border = thin
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        if title:
            ws.cell(len(plan) + 3, 1, f"任务：{title}")
        # Data sheet first
        if wb.sheetnames and wb.sheetnames[0] == sheet_name and len(wb.sheetnames) > 1:
            wb.move_sheet(wb.sheetnames[1], offset=-1)
        data_ws = wb[wb.sheetnames[0]] if wb.sheetnames else None
        if data_ws is not None and data_ws.max_row >= 1:
            try:
                data_ws.auto_filter.ref = data_ws.dimensions
                data_ws.freeze_panes = "C2" if data_ws.max_column >= 2 else "A2"
            except Exception:
                pass
        wb.save(dest)
        wb.close()
    except Exception:
        # Still return data file even if polish fails
        _append_koujing_sheet(dest, column_plan, title=title or preferred_name or run_id)
        _style_data_sheet(dest)
    return rel


def write_query_result_deliverable(
    sandbox_id: str,
    rows: list[dict[str, Any]] | None,
    *,
    column_headers: list[str],
    column_plan: list[dict] | None = None,
    preferred_name: str = "",
    title: str = "",
) -> str | None:
    """Write one CTE query result directly to a dual-sheet xlsx."""
    headers = [str(h).strip() for h in column_headers if str(h).strip()]
    if not headers:
        return None
    root = ensure_workplace(sandbox_id)
    name = preferred_name or (f"{_safe_stem(title)}.xlsx" if title else "export.xlsx")
    if not name.lower().endswith(".xlsx"):
        name += ".xlsx"
    dest = root / name
    dest.parent.mkdir(parents=True, exist_ok=True)

    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "数据"
        header_fill = PatternFill("solid", fgColor="305496")
        header_font = Font(bold=True, color="FFFFFF")
        for idx, header in enumerate(headers, 1):
            cell = ws.cell(1, idx, header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        for row_idx, row in enumerate(rows or [], 2):
            item = row if isinstance(row, dict) else {}
            for col_idx, header in enumerate(headers, 1):
                ws.cell(row_idx, col_idx, item.get(header))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for idx, header in enumerate(headers, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = max(
                12, min(40, len(header) * 2 + 4),
            )
        wb.save(dest)
        wb.close()
        _append_koujing_sheet(dest, column_plan, title=title)
        _style_data_sheet(dest)
    except Exception:
        return None
    return dest.relative_to(root).as_posix()


def write_export_from_uid_maps(
    dest: Path,
    *,
    headers: list[str],
    rows: list[dict[str, Any]],
    column_plan: list[dict] | None = None,
    title: str = "",
) -> bool:
    """Write dual-sheet xlsx from already-merged row dicts (join_on keys as columns)."""
    if not headers or not rows:
        return False
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    except ImportError:
        return False
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "数据"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    thin = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin
    for ri, row in enumerate(rows, 2):
        for ci, h in enumerate(headers, 1):
            stem = re.split(r"[（(]", h, maxsplit=1)[0].strip()
            val = row.get(h, row.get(stem, ""))
            cell = ws.cell(ri, ci, val)
            cell.border = thin
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "C2"
    wb.save(dest)
    wb.close()
    _append_koujing_sheet(dest, column_plan, title=title)
    return dest.is_file()
