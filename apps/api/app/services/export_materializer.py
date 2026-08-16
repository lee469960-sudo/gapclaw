"""Export materialization helpers for platform-written deliverables."""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from app.services.export_build_report import write_export_deliverable

logger = logging.getLogger(__name__)


@dataclass
class ExportMaterializationResult:
    file_rel: str = ""
    column_plan: list[dict[str, Any]] = field(default_factory=list)
    headers: list[str] = field(default_factory=list)
    title: str = ""


def _header_stem(h: str) -> str:
    return re.split(r"[（(]", str(h or ""), maxsplit=1)[0].strip().lower()


def _is_uid_header(h: str) -> bool:
    stem = _header_stem(h)
    n = "".join(ch for ch in stem if not ch.isspace())
    return n in ("uid", "userid", "user_id") or "用户id" in n or n.endswith("用户编号")


def _uid_col_index(headers: list[str]) -> int:
    for i, h in enumerate(headers):
        if _is_uid_header(str(h or "")):
            return i
    return -1


def merge_focus_into_prior_deliverable(
    *,
    sandbox_id: str,
    run_id: str,
    prior_rel: str,
    focus_columns: list[str] | None,
    column_plan: list[dict] | None = None,
    time_window: dict | None = None,
    title: str = "",
) -> str | None:
    """Materialize focus cols from current run pages, merge by uid into prior xlsx.

    Soft contract: on any failure return None (caller may create a new file).
    Prior full header set is preserved; only focus columns are overwritten.
    """
    from app.services.workplace import download_path, ensure_workplace, read_deliverable_headers

    sid = str(sandbox_id or "").strip()
    rid = str(run_id or "").strip()
    prior = str(prior_rel or "").strip().replace("\\", "/")
    focus = [str(c).strip() for c in (focus_columns or []) if str(c).strip()]
    if not sid or not rid or not prior or not focus:
        return None
    prior_path = download_path(sid, prior)
    if not prior_path or not prior_path.is_file():
        return None

    prior_headers = read_deliverable_headers(prior_path)
    if not prior_headers:
        return None
    # Ensure identity present for join when materializing focus patch
    fetch_headers = list(focus)
    for h in prior_headers:
        if _is_uid_header(h) and h not in fetch_headers:
            fetch_headers.insert(0, h)
            break
    else:
        # prior has no uid-like header — cannot merge safely
        return None

    temp_name = f"_fill_tmp_{rid}.xlsx"
    try:
        temp_rel = write_export_deliverable(
            sid,
            rid,
            column_plan=column_plan,
            time_window=time_window,
            column_headers=fetch_headers,
            preferred_name=temp_name,
            title=title or "列补齐",
        )
    except Exception:
        logger.exception("fill temp materialize failed")
        return None
    if not temp_rel:
        return None
    temp_path = download_path(sid, temp_rel)
    if not temp_path or not temp_path.is_file():
        return None

    try:
        import openpyxl

        wb_tmp = openpyxl.load_workbook(temp_path, read_only=True, data_only=True)
        try:
            ws_tmp = wb_tmp.active
            if ws_tmp is None:
                return None
            tmp_iter = ws_tmp.iter_rows(values_only=True)
            tmp_hdrs = [str(c).strip() if c is not None else "" for c in next(tmp_iter, ())]
            tmp_uid_i = _uid_col_index(tmp_hdrs)
            if tmp_uid_i < 0:
                return None
            tmp_focus_i: dict[str, int] = {}
            for fc in focus:
                for j, h in enumerate(tmp_hdrs):
                    if h == fc or _header_stem(h) == _header_stem(fc):
                        tmp_focus_i[fc] = j
                        break
            if not tmp_focus_i:
                return None
            patch: dict[str, dict[str, Any]] = {}
            for row in tmp_iter:
                if not row:
                    continue
                uid = str(row[tmp_uid_i] if tmp_uid_i < len(row) else "").strip()
                if not uid:
                    continue
                vals: dict[str, Any] = {}
                for fc, j in tmp_focus_i.items():
                    if j < len(row) and row[j] is not None and str(row[j]).strip() != "":
                        vals[fc] = row[j]
                if vals:
                    patch[uid] = vals
        finally:
            wb_tmp.close()

        if not patch:
            return None

        wb_prior = openpyxl.load_workbook(prior_path)
        try:
            ws_prior = wb_prior.active
            if ws_prior is None:
                return None
            prior_hdrs = [
                str(c.value).strip() if c.value is not None else ""
                for c in next(ws_prior.iter_rows(min_row=1, max_row=1))
            ]
            uid_i = _uid_col_index(prior_hdrs)
            if uid_i < 0:
                return None
            focus_idx: dict[str, int] = {}
            for fc in focus:
                for j, h in enumerate(prior_hdrs):
                    if h == fc or _header_stem(h) == _header_stem(fc):
                        focus_idx[fc] = j
                        break
            if not focus_idx:
                return None
            updated = 0
            for row in ws_prior.iter_rows(min_row=2):
                if uid_i >= len(row):
                    continue
                uid = str(row[uid_i].value or "").strip()
                if not uid or uid not in patch:
                    continue
                vals = patch[uid]
                for fc, j in focus_idx.items():
                    if fc not in vals or j >= len(row):
                        continue
                    row[j].value = vals[fc]
                    updated += 1
            if updated <= 0:
                return None
            wb_prior.save(prior_path)
        finally:
            wb_prior.close()
    except Exception:
        logger.exception("merge focus into prior failed prior=%s", prior)
        return None
    finally:
        try:
            root = ensure_workplace(sid)
            tp = root / temp_name
            if tp.is_file():
                tp.unlink()
        except Exception:
            pass

    return prior


def mark_incomplete_columns(
    column_plan: list[dict] | None,
    incomplete_roles: set[str] | list[str] | tuple[str, ...] | None,
    *,
    mark_note: bool = True,
) -> list[dict[str, Any]]:
    """Return a copy of column_plan with role-owned columns marked 未完整."""
    roles_incomplete = {
        str(r).strip().lower()
        for r in (incomplete_roles or [])
        if str(r).strip()
    }
    plan: list[dict[str, Any]] = []
    for col in column_plan or []:
        if not isinstance(col, dict):
            continue
        c = dict(col)
        roles = {
            str(s.get("role") or "").strip().lower()
            for s in (c.get("sources") or [])
            if isinstance(s, dict)
        }
        spec_role = str((c.get("agg_spec") or {}).get("role") or "").strip().lower()
        if spec_role:
            roles.add(spec_role)
        if roles & roles_incomplete:
            method = str(c.get("统计方法") or c.get("compute") or "")
            if "未完整" not in method:
                c["统计方法"] = f"{method}（未完整）" if method else "未完整"
            if mark_note:
                note = str(c.get("口径") or "")
                if "未完整" not in note:
                    c["口径"] = f"{note}；未完整" if note else "未完整"
        plan.append(c)
    return plan


def materialize_platform_export(
    *,
    sandbox_id: str,
    run_id: str,
    column_plan: list[dict] | None,
    time_window: dict | None,
    headers: list[str] | None = None,
    title: str = "",
    incomplete_roles: set[str] | list[str] | tuple[str, ...] | None = None,
    preferred_name: str = "",
    mark_note: bool = True,
    fill_prior_rel: str = "",
    focus_columns: list[str] | None = None,
) -> ExportMaterializationResult:
    plan_for_write = mark_incomplete_columns(
        column_plan,
        incomplete_roles,
        mark_note=mark_note,
    )
    out_headers = [
        str(h).strip()
        for h in (headers or [])
        if str(h).strip()
    ]
    if not out_headers:
        out_headers = [
            str(c.get("header") or "").strip()
            for c in plan_for_write
            if str(c.get("header") or "").strip()
        ]
    fill_prior = str(fill_prior_rel or "").strip().replace("\\", "/")
    focus = [str(c).strip() for c in (focus_columns or []) if str(c).strip()]
    if fill_prior and focus:
        merged = merge_focus_into_prior_deliverable(
            sandbox_id=sandbox_id,
            run_id=run_id,
            prior_rel=fill_prior,
            focus_columns=focus,
            column_plan=plan_for_write or None,
            time_window=time_window,
            title=title or "导出数据",
        )
        if merged:
            from app.services.workplace import download_path, read_deliverable_headers

            p = download_path(sandbox_id, merged)
            hdrs = read_deliverable_headers(p) if p else out_headers
            return ExportMaterializationResult(
                file_rel=merged,
                column_plan=plan_for_write,
                headers=hdrs or out_headers,
                title=title or "导出数据",
            )
        # Soft fallback: create new file below
    name = preferred_name or ""
    if fill_prior and not name:
        name = Path(fill_prior).name
    rel = write_export_deliverable(
        sandbox_id,
        run_id,
        column_plan=plan_for_write or None,
        time_window=time_window,
        column_headers=out_headers or None,
        preferred_name=name,
        title=title or "导出数据",
    )
    return ExportMaterializationResult(
        file_rel=rel or "",
        column_plan=plan_for_write,
        headers=out_headers,
        title=title or "导出数据",
    )
