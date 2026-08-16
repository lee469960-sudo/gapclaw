"""列补齐：覆盖 prior 交付 xlsx，不另写窄表；无硬门禁。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.services.export_column_plan import build_column_plan
from app.services.export_fill_scope import (
    parse_fill_scope_payload,
    resolve_export_fill_target_rel,
)
from app.services.export_materializer import (
    merge_focus_into_prior_deliverable,
    materialize_platform_export,
)
from app.services.workplace import ensure_workplace


def _write_wide(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(["用户ID", "总充值金额", "总提现金额", "总投注金额"])
    ws.append(["u1", 100, 10, ""])
    ws.append(["u2", 200, 20, ""])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _write_focus_patch(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(["用户ID", "总投注金额"])
    ws.append(["u1", 999])
    ws.append(["u2", 888])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def test_resolve_fill_target_from_prior_deliverable(tmp_path, monkeypatch):
    from app.services import workplace as wp

    sid = "sbx-fill-tgt"
    monkeypatch.setattr(wp, "_wp_root", lambda _s: tmp_path / _s / "workplace")
    root = ensure_workplace(sid)
    prior = root / "充值用户分析_wide.xlsx"
    _write_wide(prior)
    rel = resolve_export_fill_target_rel(
        sid,
        {"deliverable": "充值用户分析_wide.xlsx"},
    )
    assert rel == "充值用户分析_wide.xlsx"


def test_resolve_fill_target_soft_empty_when_missing(tmp_path, monkeypatch):
    from app.services import workplace as wp

    sid = "sbx-fill-miss"
    monkeypatch.setattr(wp, "_wp_root", lambda _s: tmp_path / _s / "workplace")
    ensure_workplace(sid)
    rel = resolve_export_fill_target_rel(
        sid,
        {"deliverable": "gone.xlsx"},
    )
    assert rel == ""


def test_merge_focus_overwrites_same_path_keeps_headers(tmp_path, monkeypatch):
    from app.services import workplace as wp
    import app.services.export_materializer as mat

    sid = "sbx-fill-merge"
    monkeypatch.setattr(wp, "_wp_root", lambda _s: tmp_path / _s / "workplace")
    root = ensure_workplace(sid)
    prior_rel = "充值用户分析_wide.xlsx"
    _write_wide(root / prior_rel)

    def _fake_write(sandbox_id, run_id, **kwargs):
        name = kwargs.get("preferred_name") or f"_fill_tmp_{run_id}.xlsx"
        _write_focus_patch(root / name)
        return name

    monkeypatch.setattr(mat, "write_export_deliverable", _fake_write)

    out = merge_focus_into_prior_deliverable(
        sandbox_id=sid,
        run_id="rid1",
        prior_rel=prior_rel,
        focus_columns=["总投注金额"],
    )
    assert out == prior_rel
    assert (root / prior_rel).exists()
    assert not (root / "_fill_tmp_rid1.xlsx").exists()

    wb = load_workbook(root / prior_rel)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert headers == ["用户ID", "总充值金额", "总提现金额", "总投注金额"]
    assert ws[2][3].value == 999
    assert ws[2][1].value == 100  # prior 充值列保留
    wb.close()


def test_materialize_platform_export_fill_prior(tmp_path, monkeypatch):
    from app.services import workplace as wp
    import app.services.export_materializer as mat

    sid = "sbx-fill-mat"
    monkeypatch.setattr(wp, "_wp_root", lambda _s: tmp_path / _s / "workplace")
    root = ensure_workplace(sid)
    prior_rel = "接口表.xlsx"
    _write_wide(root / prior_rel)

    def _fake_write(sandbox_id, run_id, **kwargs):
        name = kwargs.get("preferred_name") or f"_fill_tmp_{run_id}.xlsx"
        _write_focus_patch(root / name)
        return name

    monkeypatch.setattr(mat, "write_export_deliverable", _fake_write)

    plan = build_column_plan(["用户ID", "总投注金额"])
    result = materialize_platform_export(
        sandbox_id=sid,
        run_id="rid",
        column_plan=plan,
        time_window={"start_ms": 1, "end_ms": 2, "label": "fill"},
        headers=["用户ID", "总充值金额", "总提现金额", "总投注金额"],
        preferred_name="用户分析_rid.xlsx",
        fill_prior_rel=prior_rel,
        focus_columns=["总投注金额"],
    )
    assert result.file_rel == prior_rel
    assert result.file_rel != "用户分析_rid.xlsx"
    assert len(result.headers) == 4
    assert "总投注金额" in result.headers


def test_materialize_soft_new_file_when_no_prior(tmp_path, monkeypatch):
    from app.services import workplace as wp
    import app.services.export_materializer as mat

    sid = "sbx-fill-soft"
    monkeypatch.setattr(wp, "_wp_root", lambda _s: tmp_path / _s / "workplace")
    root = ensure_workplace(sid)

    def _fake_write(sandbox_id, run_id, **kwargs):
        name = kwargs.get("preferred_name") or "用户分析_new.xlsx"
        _write_focus_patch(root / name)
        return name

    monkeypatch.setattr(mat, "write_export_deliverable", _fake_write)

    result = materialize_platform_export(
        sandbox_id=sid,
        run_id="rid2",
        column_plan=build_column_plan(["用户ID", "总投注金额"]),
        time_window=None,
        headers=["用户ID", "总投注金额"],
        preferred_name="用户分析_rid2.xlsx",
        fill_prior_rel="",
        focus_columns=["总投注金额"],
    )
    assert result.file_rel == "用户分析_rid2.xlsx"
    assert (root / result.file_rel).is_file()


def test_fill_scope_no_role_invention_from_utterance():
    scope = parse_fill_scope_payload(
        {
            "is_column_fill": True,
            "focus_columns": ["总投注金额"],
            "include_identity": True,
            "reason": "补投注",
        },
        prior_headers=["用户ID", "总充值金额", "总投注金额"],
    )
    assert scope.is_column_fill
    assert scope.focus_columns == ["总投注金额"]
