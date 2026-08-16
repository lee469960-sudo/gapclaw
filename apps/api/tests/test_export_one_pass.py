"""Generic one-pass export: query graph cost/segment + platform dual-sheet write."""

import json
from pathlib import Path

from app.services.export_build_report import (
    write_export_deliverable,
    write_export_from_uid_maps,
)
from app.services.export_column_plan import (
    FETCH_AGG,
    FETCH_SEQUENCE,
    FETCH_TOP_N,
    build_column_plan,
    build_query_graph,
    node_cost_for_mode,
    node_segment_for,
    sql_from_agg_spec,
)
from app.services.workplace import ensure_workplace, write_task_json_page


_GOLDEN_16 = [
    "注册时间",
    "用户ID",
    "注册渠道",
    "总充值金额",
    "总提现金额",
    "当前余额(SC)",
    "总下注金额(SC)",
    "总返奖金额(SC)",
    "下注次数",
    "流水倍数",
    "连续充值次数",
    "SC投注金额最多的游戏",
    "是否被封禁",
    "是否有退款",
    "充值银行卡数量",
    "提现银行卡数量",
]


def test_pdf_koujing_keywords():
    plan = build_column_plan(_GOLDEN_16)
    by_h = {c["header"]: c for c in plan}
    pay = by_h["总充值金额"]
    assert "finish_time" in (pay.get("统计方法") or "") or pay["agg_spec"].get("time_field") == "finish_time"
    assert "price" in (pay.get("统计方法") or "") or "price" in pay["agg_spec"].get("select", "")
    cash = by_h["总提现金额"]
    assert cash["agg_spec"].get("time_field") == "update_time"
    bet = by_h["总下注金额(SC)"]
    assert "everyday_bygame" in bet["agg_spec"].get("view", "")
    assert "goods_type" in (bet.get("统计方法") or "")
    cards = by_h["充值银行卡数量"]
    assert "channel_id=11" in (cards.get("统计方法") or "")
    assert "cardNumber" in (cards.get("统计方法") or "")
    assert "cardNumber" in (cards["agg_spec"].get("select") or "")
    assert "count(DISTINCT account)" not in (cards["agg_spec"].get("select") or "")
    cash_cards = by_h["提现银行卡数量"]
    assert "disbursementNumber" in (cash_cards.get("统计方法") or "")
    assert "disbursementNumber" in (cash_cards["agg_spec"].get("select") or "")
    assert cash_cards["agg_spec"].get("time_field") == "update_time"


def test_query_graph_has_cost_and_segment():
    plan = build_column_plan([
        "用户ID",
        "总充值金额",
        "总下注金额(SC)",
        "流水倍数",
        "连续充值次数",
        "SC投注金额最多的游戏",
    ])
    tw = {"start_ms": 1, "end_ms": 2, "cohort": "pay"}
    nodes = build_query_graph(plan, tw)
    assert nodes
    assert next(n for n in nodes if n.get("cost") == "light")["role"] == "pay"
    assert next(n for n in nodes if n.get("role") == "user").get("deferred") is True
    light_roles = [n.get("role") for n in nodes if n.get("cost") == "light"]
    assert light_roles.index("user") < light_roles.index("bet")
    user_idx = next(i for i, n in enumerate(nodes) if n.get("role") == "user")
    secondary_pay_idx = [
        i for i, n in enumerate(nodes)
        if (
            n.get("role") == "pay"
            and n.get("segment") != "cohort"
            and "pay_sum" not in str(n.get("sql") or "")
        )
    ]
    assert secondary_pay_idx and user_idx < min(secondary_pay_idx)
    assert all("cost" in n and "segment" in n for n in nodes)
    costs = {n["cost"] for n in nodes}
    assert "light" in costs
    # heavy nodes last
    assert node_cost_for_mode(FETCH_SEQUENCE) == "heavy"
    assert node_cost_for_mode(FETCH_TOP_N) == "heavy"
    assert node_cost_for_mode(FETCH_AGG) == "light"
    assert node_segment_for(FETCH_AGG, "pay", "view_result_pay_order_log") == "fact_agg"
    assert "bet_daily" in node_segment_for(
        FETCH_AGG, "bet", "view_result_gameuser_betstat_everyday_bygame"
    )


def test_agg_sql_uses_finish_time_and_price():
    plan = build_column_plan(["总充值金额"])
    view, sql = sql_from_agg_spec(plan[0]["agg_spec"], {"start_ms": 1000, "end_ms": 2000})
    assert view == "view_result_pay_order_log"
    assert "finish_time" in sql
    assert "sum(price)" in sql.lower()
    assert "status = 2" in sql


def test_refund_amount_materializes_separately_from_refund_flag(tmp_path, monkeypatch):
    sid = "sbx_refund_amount"
    root = tmp_path / sid / "workplace"
    root.mkdir(parents=True)

    def _wp(s):
        return root if s == sid else Path("/nope")

    monkeypatch.setattr("app.services.workplace.ensure_workplace", _wp)
    monkeypatch.setattr("app.services.export_build_report.ensure_workplace", _wp)

    run_id = "run_refund_amount"
    write_task_json_page(
        sid,
        [{"uid": 1, "register_time": 1500}],
        run_id=run_id,
        page=1,
        view="view_result_user_info",
    )
    write_task_json_page(
        sid,
        [
            {"uid": 1, "status": 2, "price": 1000, "create_time": 1500},
            {"uid": 1, "status": 4, "price": 225, "update_time": 1600},
        ],
        run_id=run_id,
        page=2,
        view="view_result_pay_order_log",
    )
    plan = build_column_plan(["用户ID", "退款金额", "是否有退款"])
    rel = write_export_deliverable(
        sid,
        run_id,
        column_plan=plan,
        time_window={"start_ms": 1000, "end_ms": 2000, "label": "测试窗"},
        title="退款测试",
    )
    assert rel

    import openpyxl

    wb = openpyxl.load_workbook(root / rel, data_only=True)
    try:
        row = list(wb["数据"].iter_rows(min_row=2, max_row=2, values_only=True))[0]
    finally:
        wb.close()
    assert row == ("1", 2.25, "有")


def test_generic_small_plan_zero_shell_write(tmp_path, monkeypatch):
    """3-col plan + fake task pages → dual-sheet xlsx without LLM SHELL."""
    sid = "sbx_onepass"
    root = tmp_path / sid / "workplace"
    root.mkdir(parents=True)

    def _wp(s):
        return root if s == sid else Path("/nope")

    monkeypatch.setattr("app.services.workplace.ensure_workplace", _wp)
    monkeypatch.setattr("app.services.export_build_report.ensure_workplace", _wp)

    run_id = "run1"
    users = [
        {"uid": 1, "register_time": 1500, "channel_id": 9, "sc": 10000},
        {"uid": 2, "register_time": 1600, "channel_id": 9, "sc": 0},
    ]
    pays = [{"uid": 1, "pay_sum": 12.5}, {"uid": 2, "pay_sum": 3.0}]
    channels = [{"channel_id": 9, "channel_name": "Organic"}]
    write_task_json_page(sid, users, run_id=run_id, page=1, view="view_result_user_info")
    write_task_json_page(sid, pays, run_id=run_id, page=2, view="view_result_pay_order_log")
    write_task_json_page(sid, channels, run_id=run_id, page=3, view="view_result_config_channel")

    plan = build_column_plan(["用户ID", "总充值金额", "注册渠道"])
    rel = write_export_deliverable(
        sid,
        run_id,
        column_plan=plan,
        time_window={"start_ms": 1000, "end_ms": 2000, "label": "测试窗", "cohort": "pay"},
        title="测试导出",
    )
    assert rel
    dest = root / rel
    assert dest.is_file()
    import openpyxl
    wb = openpyxl.load_workbook(dest)
    assert "口径说明" in wb.sheetnames
    assert any(n in ("数据", "Sheet1", "Sheet") or "分析" in n for n in wb.sheetnames)
    notes = wb["口径说明"]
    blob = " ".join(
        str(v or "") for row in notes.iter_rows(values_only=True) for v in row
    )
    assert "用户ID" in blob or "总充值" in blob
    wb.close()


def test_pay_cohort_materializer_keeps_users_registered_before_window(tmp_path, monkeypatch):
    sid = "sbx_pay_cohort_materialize"
    root = tmp_path / sid / "workplace"
    root.mkdir(parents=True)

    def _wp(s):
        return root if s == sid else Path("/nope")

    monkeypatch.setattr("app.services.workplace.ensure_workplace", _wp)
    monkeypatch.setattr("app.services.export_build_report.ensure_workplace", _wp)

    run_id = "run_pay"
    headers = [
        "注册时间",
        "用户ID",
        "注册渠道",
        "总充值金额",
        "总提现金额",
        "当前余额(SC)",
        "总下注金额(SC)",
        "总返奖金额(SC)",
        "下注次数(SC投注次数)",
        "流水倍数(时间段内总下注金额/总充值金额)",
        "连续充值次数(两次游戏行为之间的充值次数，需要最多的次数)",
        "SC投注金额最多的游戏(该用户在游戏中SC投注最多的游戏，格式：游戏名称，游戏ID)",
        "是否被封禁(封禁类型：充值/提现/登录，根据封禁类型显示)",
        "是否有退款(是否有退款成功的订单，有则显示文本\"有\"，无则空)",
        "充值银行卡数量(APTPAY充值渠道使用的银行卡数量，需排重)",
        "提现银行卡数量(APTPAY提现渠道使用的银行卡数量，需排重)",
    ]
    write_task_json_page(
        sid,
        [{"uid": 7, "register_time": 500, "channel_id": 1, "sc": 1234}],
        run_id=run_id,
        page=1,
        view="view_result_user_info",
    )
    write_task_json_page(
        sid,
        [{"uid": 7, "pay_sum": 20.0}],
        run_id=run_id,
        page=2,
        view="view_result_pay_order_log",
    )
    write_task_json_page(
        sid,
        [{"channel_id": 1, "channel_name": "Ads"}],
        run_id=run_id,
        page=3,
        view="view_result_config_channel",
    )
    write_task_json_page(
        sid,
        [{"uid": 7, "card_cnt": 2}],
        run_id=run_id,
        page=4,
        view="view_result_pay_order_log",
    )
    write_task_json_page(
        sid,
        [{"uid": 7, "card_cnt": 1}],
        run_id=run_id,
        page=5,
        view="view_result_cash_order_log",
    )
    write_task_json_page(
        sid,
        [
            {"uid": 7, "create_time": 1200, "status": 2},
            {"uid": 7, "create_time": 1300, "status": 2},
            {"uid": 7, "create_time": 1500, "status": 2},
        ],
        run_id=run_id,
        page=6,
        view="view_result_pay_order_log",
    )
    write_task_json_page(
        sid,
        [
            {"uid": 7, "create_time": 1100, "type": 1, "game_id": 9, "bet_sc": 1},
            {"uid": 7, "create_time": 1400, "type": 1, "game_id": 9, "bet_sc": 1},
            {"uid": 7, "create_time": 1600, "type": 1, "game_id": 9, "bet_sc": 1},
        ],
        run_id=run_id,
        page=7,
        view="view_result_user_bet_log",
    )
    write_task_json_page(
        sid,
        [{"game_id": 10, "game_name": "Lucky Spin"}],
        run_id=run_id,
        page=8,
        view="view_result_config_game",
    )
    write_task_json_page(
        sid,
        [{"uid": 7, "game_id": 10, "amt": 9.0}],
        run_id=run_id,
        page=9,
        view="view_result_user_bet_log",
    )

    rel = write_export_deliverable(
        sid,
        run_id,
        column_plan=build_column_plan(headers),
        time_window={"start_ms": 1000, "end_ms": 2000, "label": "pay窗", "cohort": "pay"},
        title="pay cohort",
    )
    assert rel
    import openpyxl
    wb = openpyxl.load_workbook(root / rel)
    ws = wb["数据"]
    assert ws.max_row == 2
    out_headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    assert out_headers == headers
    row = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
    assert row[out_headers.index("用户ID")] == "7"
    assert row[out_headers.index("总充值金额")] == 20
    assert row[out_headers.index("注册渠道")] == "Ads"
    assert row[out_headers.index("连续充值次数(两次游戏行为之间的充值次数，需要最多的次数)")] == 2
    assert row[out_headers.index("SC投注金额最多的游戏(该用户在游戏中SC投注最多的游戏，格式：游戏名称，游戏ID)")] == "Lucky Spin,10"
    assert row[out_headers.index("充值银行卡数量(APTPAY充值渠道使用的银行卡数量，需排重)")] == 2
    assert row[out_headers.index("提现银行卡数量(APTPAY提现渠道使用的银行卡数量，需排重)")] == 1
    wb.close()


def test_pay_cohort_materializer_keeps_paid_users_without_user_info(tmp_path, monkeypatch):
    sid = "sbx_pay_cohort_missing_userinfo"
    root = tmp_path / sid / "workplace"
    root.mkdir(parents=True)

    def _wp(s):
        return root if s == sid else Path("/nope")

    monkeypatch.setattr("app.services.workplace.ensure_workplace", _wp)
    monkeypatch.setattr("app.services.export_build_report.ensure_workplace", _wp)

    run_id = "run_pay_partial_user"
    headers = ["用户ID", "总充值金额", "注册渠道"]
    write_task_json_page(
        sid,
        [{"uid": 7, "register_time": 500, "channel_id": 1}],
        run_id=run_id,
        page=1,
        view="view_result_user_info",
    )
    write_task_json_page(
        sid,
        [{"uid": 7, "pay_sum": 20.0}, {"uid": 8, "pay_sum": 30.0}],
        run_id=run_id,
        page=2,
        view="view_result_pay_order_log",
    )
    write_task_json_page(
        sid,
        [{"channel_id": 1, "channel_name": "Ads"}],
        run_id=run_id,
        page=3,
        view="view_result_config_channel",
    )

    rel = write_export_deliverable(
        sid,
        run_id,
        column_plan=build_column_plan(headers),
        time_window={"start_ms": 1000, "end_ms": 2000, "label": "pay窗", "cohort": "pay"},
        title="pay partial user",
    )
    assert rel

    import openpyxl

    wb = openpyxl.load_workbook(root / rel)
    ws = wb["数据"]
    out_headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    rows = [
        [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        for r in range(2, ws.max_row + 1)
    ]
    by_uid = {str(row[out_headers.index("用户ID")]): row for row in rows}
    assert set(by_uid) == {"7", "8"}
    assert by_uid["7"][out_headers.index("总充值金额")] == 20
    assert by_uid["7"][out_headers.index("注册渠道")] == "Ads"
    assert by_uid["8"][out_headers.index("总充值金额")] == 30
    assert by_uid["8"][out_headers.index("注册渠道")] in ("", None)
    wb.close()


def test_pay_cohort_materializer_uses_cohort_uid_anchor_rows(tmp_path, monkeypatch):
    sid = "sbx_pay_cohort_uid_anchor"
    root = tmp_path / sid / "workplace"
    root.mkdir(parents=True)

    def _wp(s):
        return root if s == sid else Path("/nope")

    monkeypatch.setattr("app.services.workplace.ensure_workplace", _wp)
    monkeypatch.setattr("app.services.export_build_report.ensure_workplace", _wp)

    run_id = "run_pay_uid_anchor"
    headers = ["用户ID", "总充值金额", "注册渠道"]
    write_task_json_page(
        sid,
        [{"uid": 7, "register_time": 500, "channel_id": 1}],
        run_id=run_id,
        page=1,
        view="view_result_user_info",
    )
    write_task_json_page(
        sid,
        [{"uid": 7, "cohort_uid": 1}, {"uid": 8, "cohort_uid": 1}],
        run_id=run_id,
        page=2,
        view="view_result_pay_order_log",
    )
    write_task_json_page(
        sid,
        [{"uid": 7, "pay_sum": 20.0}],
        run_id=run_id,
        page=3,
        view="view_result_pay_order_log",
    )
    write_task_json_page(
        sid,
        [{"channel_id": 1, "channel_name": "Ads"}],
        run_id=run_id,
        page=4,
        view="view_result_config_channel",
    )

    rel = write_export_deliverable(
        sid,
        run_id,
        column_plan=build_column_plan(headers),
        time_window={"start_ms": 1000, "end_ms": 2000, "label": "pay窗", "cohort": "pay"},
        title="pay uid anchor",
    )
    assert rel

    import openpyxl

    wb = openpyxl.load_workbook(root / rel)
    ws = wb["数据"]
    out_headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    rows = [
        [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        for r in range(2, ws.max_row + 1)
    ]
    by_uid = {str(row[out_headers.index("用户ID")]): row for row in rows}
    assert set(by_uid) == {"7", "8"}
    assert by_uid["7"][out_headers.index("总充值金额")] == 20
    assert by_uid["8"][out_headers.index("总充值金额")] in (0, None)
    assert by_uid["8"][out_headers.index("注册渠道")] in ("", None)
    wb.close()


def test_materializer_merges_all_bet_fact_views(tmp_path, monkeypatch):
    sid = "sbx_bet_fact_merge"
    root = tmp_path / sid / "workplace"
    root.mkdir(parents=True)

    def _wp(s):
        return root if s == sid else Path("/nope")

    monkeypatch.setattr("app.services.workplace.ensure_workplace", _wp)
    monkeypatch.setattr("app.services.export_build_report.ensure_workplace", _wp)

    run_id = "run_bet_merge"
    headers = [
        "用户ID",
        "总充值金额",
        "总下注金额(SC)",
        "总返奖金额(SC)",
        "下注次数(SC投注次数)",
        "流水倍数(时间段内总下注金额/总充值金额)",
    ]
    write_task_json_page(
        sid,
        [{"uid": 7, "register_time": 500}],
        run_id=run_id,
        page=1,
        view="view_result_user_info",
    )
    write_task_json_page(
        sid,
        [{"uid": 7, "pay_sum": 20.0}],
        run_id=run_id,
        page=2,
        view="view_result_pay_order_log",
    )
    write_task_json_page(
        sid,
        [{"uid": 7, "create_time": 1100}],
        run_id=run_id,
        page=3,
        view="view_result_user_bet_log",
    )
    write_task_json_page(
        sid,
        [{"uid": 7, "bet_sum": 12.0}],
        run_id=run_id,
        page=4,
        view="view_result_gameuser_betstat_everyday_bygame",
    )
    write_task_json_page(
        sid,
        [{"uid": 7, "win_sum": 5.0}],
        run_id=run_id,
        page=5,
        view="view_result_gameuser_betstat_everyday_bygame",
    )
    write_task_json_page(
        sid,
        [{"uid": 7, "bet_cnt": 3}],
        run_id=run_id,
        page=6,
        view="view_result_gameuser_betstat_everyday_bygame",
    )

    rel = write_export_deliverable(
        sid,
        run_id,
        column_plan=build_column_plan(headers),
        time_window={"start_ms": 1000, "end_ms": 2000, "label": "pay窗", "cohort": "pay"},
        title="bet merge",
    )
    assert rel
    import openpyxl

    wb = openpyxl.load_workbook(root / rel)
    ws = wb["数据"]
    out_headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    row = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
    assert row[out_headers.index("用户ID")] == "7"
    assert row[out_headers.index("总下注金额(SC)")] == 12
    assert row[out_headers.index("总返奖金额(SC)")] == 5
    assert row[out_headers.index("下注次数(SC投注次数)")] == 3
    assert row[out_headers.index("流水倍数(时间段内总下注金额/总充值金额)")] == 0.6
    wb.close()


def test_write_from_uid_maps_dual_sheet(tmp_path):
    dest = tmp_path / "out.xlsx"
    plan = build_column_plan(["用户ID", "总充值金额"])
    ok = write_export_from_uid_maps(
        dest,
        headers=["用户ID", "总充值金额"],
        rows=[{"用户ID": 1, "总充值金额": 9.9}],
        column_plan=plan,
        title="小表",
    )
    assert ok
    import openpyxl
    wb = openpyxl.load_workbook(dest)
    assert "口径说明" in wb.sheetnames
    wb.close()
