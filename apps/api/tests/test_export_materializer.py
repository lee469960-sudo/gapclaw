"""Platform export materializer contract."""

from pathlib import Path

from app.services.export_column_plan import build_column_plan
from app.services.export_materializer import (
    mark_incomplete_columns,
    materialize_platform_export,
)
from app.services.workplace import write_task_json_page


def test_mark_incomplete_columns_copies_and_marks_role():
    plan = build_column_plan(["用户ID", "总充值金额", "总提现金额"])
    marked = mark_incomplete_columns(plan, {"cash"})
    assert marked is not plan
    cash_col = next(c for c in marked if c.get("header") == "总提现金额")
    assert "未完整" in str(cash_col.get("统计方法") or "")
    assert "未完整" in str(cash_col.get("口径") or "")
    original_cash = next(c for c in plan if c.get("header") == "总提现金额")
    assert "未完整" not in str(original_cash.get("统计方法") or "")


def test_materialize_platform_export_writes_marked_deliverable(tmp_path, monkeypatch):
    sid = "sbx_materializer"
    root = tmp_path / sid / "workplace"
    root.mkdir(parents=True)

    def _wp(sandbox_id: str) -> Path:
        return root if sandbox_id == sid else Path("/nope")

    monkeypatch.setattr("app.services.workplace.ensure_workplace", _wp)
    monkeypatch.setattr("app.services.export_build_report.ensure_workplace", _wp)

    run_id = "run_materializer"
    write_task_json_page(
        sid,
        [{"uid": 1, "register_time": 1500, "sc": 10000}],
        run_id=run_id,
        page=1,
        view="view_result_user_info",
    )
    write_task_json_page(
        sid,
        [{"uid": 1, "pay_sum": 12.5}],
        run_id=run_id,
        page=2,
        view="view_result_pay_order_log",
    )
    plan = build_column_plan(["用户ID", "总充值金额", "总提现金额"])
    result = materialize_platform_export(
        sandbox_id=sid,
        run_id=run_id,
        column_plan=plan,
        time_window={"start_ms": 1000, "end_ms": 2000, "label": "mat"},
        title="materializer",
        incomplete_roles={"cash"},
    )
    assert result.file_rel
    assert (root / result.file_rel).is_file()
    assert any(
        c.get("header") == "总提现金额" and "未完整" in str(c.get("口径") or "")
        for c in result.column_plan
    )

    import openpyxl

    wb = openpyxl.load_workbook(root / result.file_rel, read_only=True, data_only=True)
    try:
        blob = " ".join(
            str(v or "")
            for row in wb["口径说明"].iter_rows(values_only=True)
            for v in row
        )
    finally:
        wb.close()
    assert "总提现金额" in blob
    assert "未完整" in blob


def test_materializer_merges_user_aux_view_rows_for_ban_columns(tmp_path, monkeypatch):
    sid = "sbx_materializer_ban"
    root = tmp_path / sid / "workplace"
    root.mkdir(parents=True)

    def _wp(sandbox_id: str) -> Path:
        return root if sandbox_id == sid else Path("/nope")

    monkeypatch.setattr("app.services.workplace.ensure_workplace", _wp)
    monkeypatch.setattr("app.services.export_build_report.ensure_workplace", _wp)

    run_id = "run_materializer_ban"
    write_task_json_page(
        sid,
        [{"uid": 1, "register_time": 1500, "sc": 10000}],
        run_id=run_id,
        page=1,
        view="view_result_user_info",
    )
    write_task_json_page(
        sid,
        [{"uid": 1, "ban_type": 3}],
        run_id=run_id,
        page=2,
        view="view_result_user_ban_log",
    )
    plan = build_column_plan(["用户ID", "是否被封禁"])
    result = materialize_platform_export(
        sandbox_id=sid,
        run_id=run_id,
        column_plan=plan,
        time_window={"start_ms": 1000, "end_ms": 2000, "label": "mat"},
        title="materializer-ban",
        incomplete_roles=set(),
    )
    assert result.file_rel

    import openpyxl

    wb = openpyxl.load_workbook(root / result.file_rel, read_only=True, data_only=True)
    try:
        ws = wb["数据"]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    assert rows[0] == ("用户ID", "是否被封禁")
    assert rows[1][0] in ("1", 1)
    assert rows[1][1] == "充值"
