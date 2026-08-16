"""Metric registry, TaskSpec, Trace, Verifier — export intelligence layer."""

from pathlib import Path

from app.services.export_column_plan import (
    build_column_plan,
    build_query_graph,
    sql_from_agg_spec,
)
from app.services.export_repair_plan import (
    activate_repair_plan_from_prior,
    activate_repair_plan_from_trace,
    build_repair_plan,
    limit_query_nodes_for_repair,
    prioritize_query_nodes_for_repair,
    repair_plan_should_reverify_after_progress,
)
from app.services.export_engine_outcome import build_engine_outcome
from app.services.export_finalizer import verify_export_finalizer
from app.services.export_task_spec import build_task_spec
from app.services.export_task_validator import (
    format_task_spec_validation_reply,
    validate_task_spec,
)
from app.services.export_trace import (
    ExportTrace,
    classify_mcp_error,
    write_export_trace,
    read_export_trace,
)
from app.services.export_verifier import (
    build_final_summary_from_verifier,
    format_verifier_block,
    verify_export_deliverable,
)
from app.services.metric_registry import (
    COHORT_TIME_FIELD,
    GOLDEN_16_HEADERS,
    PAY_METRIC_TIME_FIELD,
    match_metric,
)
from app.services.task_policy import route_export_view_intent
from app.services.export_build_report import write_export_deliverable
from app.services.workplace import write_task_json_page


def test_golden16_registry_and_pay_koujing():
    plan = build_column_plan(list(GOLDEN_16_HEADERS))
    assert len(plan) == 16
    by_h = {c["header"]: c for c in plan}
    pay = by_h["总充值金额"]
    assert pay.get("metric_id") == "pay_sum" or "finish_time" in (pay.get("统计方法") or "")
    assert pay["agg_spec"].get("time_field") == PAY_METRIC_TIME_FIELD
    assert "sum(price)" in (pay["agg_spec"].get("select") or "").lower()
    assert "status = 2" in (pay["agg_spec"].get("filter_sql") or "")

    cards = by_h["充值银行卡数量"]
    sel = cards["agg_spec"].get("select") or ""
    assert "JSONExtractString" in sel and "cardNumber" in sel
    assert "DISTINCT account)" not in sel

    cash_cards = by_h["提现银行卡数量"]
    csel = cash_cards["agg_spec"].get("select") or ""
    assert "disbursementNumber" in csel

    bet = by_h["总下注金额(SC)"]
    assert "everyday_bygame" in (bet["agg_spec"].get("view") or "")
    assert "goods_type IN (1, 4)" in (bet["agg_spec"].get("filter_sql") or "")

    streak = by_h["连续充值次数"]
    assert streak.get("fetch_mode") == "sequence"
    assert streak.get("failure_policy") == "soft_fail" or (
        streak.get("agg_spec") or {}
    ).get("incomplete_ok")

    m = match_metric("总充值金额")
    assert m is not None
    assert m.time_field == PAY_METRIC_TIME_FIELD


def test_refund_amount_and_refund_flag_use_distinct_metrics():
    amount = match_metric("退款金额")
    flag = match_metric("是否有退款")
    assert amount is not None and amount.metric_id == "refund_amount"
    assert flag is not None and flag.metric_id == "has_refund"

    plan = build_column_plan(["退款金额", "是否有退款"])
    by_header = {col["header"]: col for col in plan}
    amount_spec = by_header["退款金额"]["agg_spec"]
    assert by_header["退款金额"]["fetch_mode"] == "agg"
    assert "sum(price)" in (amount_spec.get("select") or "").lower()
    assert amount_spec.get("filter_sql") == "status = 4"
    assert by_header["是否有退款"]["fetch_mode"] == "flag"


def test_legacy_column_aliases_are_now_metric_contracts():
    expected = {
        "充值": "pay_sum",
        "提现": "cash_sum",
        "投注": "bet_sum_sc",
        "返奖": "win_sum_sc",
        "渠道": "register_channel",
        "游戏名": "top_game_by_bet",
        "余额": "balance_sc",
        "封禁": "banned",
        "姓": "first_name",
        "名": "last_name",
        "手机": "phone",
        "邮箱": "email",
    }
    for header, metric_id in expected.items():
        metric = match_metric(header)
        assert metric is not None and metric.metric_id == metric_id
        plan = build_column_plan([header])
        assert plan[0].get("metric_id") == metric_id


def test_task_spec_leaves_metric_time_semantics_to_query_contract():
    route = route_export_view_intent(
        "导出美东时间充值用户：总充值金额、流水倍数、充值银行卡数量",
    )
    assert route.mode in ("multi_fact", "ambiguous") or "multi" in (route.mode or "")
    spec = build_task_spec(
        task_type="multi_fact",
        headers=list(GOLDEN_16_HEADERS),
        time_window={
            "start_ms": 1,
            "end_ms": 2,
            "cohort": "pay",
            "label": "test",
        },
        resources=["pay_orders", "cash_orders", "bet_daily"],
        source_brief="充值用户导出",
    )
    assert spec.cohort.type == "pay"
    assert spec.cohort.time_field == ""
    assert "metric_time_field_pay" not in spec.to_dict()
    assert "cohort_time_field" not in spec.to_dict()
    validation = validate_task_spec(spec)
    assert validation.status == "pass"
    assert validation.ok


def test_task_spec_records_bound_resources_not_legacy_roles():
    spec = build_task_spec(
        task_type="multi_fact",
        headers=["是否被封禁"],
        time_window={"start_ms": 1, "end_ms": 2},
        resources=["view_result_user_ban_status"],
    )
    assert spec.required_resources == ["view_result_user_ban_status"]
    assert "required_roles" not in spec.to_dict()


def test_type_a_light_identity_roles():
    plan = build_column_plan(["用户ID", "注册时间", "注册渠道"])
    roles = {str(s.get("role")) for c in plan for s in (c.get("sources") or [])}
    assert "user" in roles
    assert "pay" not in roles
    assert "bet" not in roles
    spec = build_task_spec(
        task_type="light_identity",
        headers=["用户ID", "注册时间"],
        resources=["user_directory"],
    )
    assert spec.task_type == "light_identity"
    assert spec.required_resources == ["user_directory"]
    validation = validate_task_spec(spec)
    assert validation.status == "clarification"
    assert any(i.code == "missing_time_window" for i in validation.issues)


def test_type_c_pinned_view_task_spec():
    spec = build_task_spec(
        task_type="single_view",
        headers=[],
        pinned_views=["view_result_pay_order_log"],
        resources=[],
    )
    assert spec.task_type == "single_view"
    assert spec.pinned_views == ["view_result_pay_order_log"]
    assert spec.cohort.type == "custom"
    validation = validate_task_spec(spec)
    assert validation.status == "pass"


def test_task_spec_validator_does_not_invent_business_time_field_rules():
    spec = build_task_spec(
        task_type="multi_fact",
        headers=["用户ID", "总充值金额"],
        time_window={
            "start_ms": 1,
            "end_ms": 2,
            "cohort": "pay",
        },
        resources=["pay_orders"],
        source_brief="充值用户导出",
    )
    spec.cohort.time_field = PAY_METRIC_TIME_FIELD
    validation = validate_task_spec(spec)
    assert validation.status == "pass"
    assert not any(i.code.startswith("invalid_pay_") for i in validation.issues)


def test_task_spec_validator_requires_single_view_pin():
    spec = build_task_spec(
        task_type="single_view",
        headers=[],
        pinned_views=[],
        resources=[],
    )
    validation = validate_task_spec(spec)
    assert validation.status == "clarification"
    assert any(i.code == "missing_pinned_view" for i in validation.issues)
    reply = format_task_spec_validation_reply(validation)
    assert "暂时不会开始拉数" in reply
    assert "宣称任务已完成" not in reply


def test_task_spec_validator_error_reply_blocks_completion_claim():
    spec = build_task_spec(
        task_type="multi_fact",
        headers=["用户ID", "总充值金额"],
        time_window={
            "start_ms": 2,
            "end_ms": 1,
            "cohort": "pay",
        },
        resources=["pay_orders"],
        source_brief="充值用户导出",
    )
    validation = validate_task_spec(spec)
    assert validation.status == "error"
    reply = format_task_spec_validation_reply(validation)
    assert "暂时不会执行导出" in reply
    assert "不会在契约不合法时宣称任务已完成" in reply


def test_export_trace_write_and_classify(tmp_path, monkeypatch):
    sid = "sbx_trace"
    root = tmp_path / sid / "workplace"
    root.mkdir(parents=True)

    def _wp(s):
        return root if s == sid else Path("/nope")

    monkeypatch.setattr("app.services.workplace.ensure_workplace", _wp)
    monkeypatch.setattr("app.services.export_trace.ensure_workplace", _wp)

    tr = ExportTrace(run_id="run1", source_brief="brief")
    tr.export_contract = {"version": "export_contract.v1", "task_spec": {"task_type": "multi_fact"}}
    tr.task_spec_validation = {"status": "pass", "ok": True}
    tr.set_repair_plan({"status": "none", "actions": []})
    tr.record_query(
        key="agg:pay",
        role="pay",
        view="view_result_pay_order_log",
        error='Unknown expression or function identifier \'amount\'',
        status="error",
    )
    assert tr.query_nodes[0]["error_class"] == "unknown_column"
    assert classify_mcp_error("timeout after 30s") == "timeout"
    rel = write_export_trace(sid, "run1", tr)
    assert rel.endswith("export_trace.json")
    loaded = read_export_trace(sid, "run1")
    assert loaded and loaded.get("run_id") == "run1"
    assert loaded.get("export_contract", {}).get("version") == "export_contract.v1"
    assert loaded.get("task_spec_validation", {}).get("status") == "pass"
    assert loaded.get("repair_plan", {}).get("status") == "none"


def test_verifier_passes_mock_deliverable(tmp_path, monkeypatch):
    sid = "sbx_vr"
    root = tmp_path / sid / "workplace"
    root.mkdir(parents=True)

    def _wp(s):
        return root if s == sid else Path("/nope")

    monkeypatch.setattr("app.services.workplace.ensure_workplace", _wp)
    monkeypatch.setattr("app.services.export_build_report.ensure_workplace", _wp)
    monkeypatch.setattr("app.services.export_verifier.ensure_workplace", _wp)

    run_id = "run_vr"
    write_task_json_page(
        sid,
        [{"uid": 1, "register_time": 1500, "channel_id": 9, "sc": 10000}],
        run_id=run_id,
        page=1,
        view="view_result_user_info",
    )
    write_task_json_page(
        sid,
        [{"uid": 1, "pay_sum": 12.5}],
        run_id=run_id,
        page=2,
        view="view_result_pay_order_log",
    )
    write_task_json_page(
        sid,
        [{"channel_id": 9, "channel_name": "Organic"}],
        run_id=run_id,
        page=3,
        view="view_result_config_channel",
    )
    plan = build_column_plan(["用户ID", "总充值金额", "注册渠道"])
    rel = write_export_deliverable(
        sid,
        run_id,
        column_plan=plan,
        time_window={"start_ms": 1000, "end_ms": 2000, "label": "vr", "cohort": "pay"},
        title="verifier",
    )
    assert rel
    tr = ExportTrace(run_id=run_id)
    tr.record_query(
        key="pay", role="pay", view="view_result_pay_order_log",
        row_count=1, status="ok",
    )
    tr.record_query(
        key="user", role="user", view="view_result_user_info",
        row_count=1, status="ok",
    )
    result = verify_export_deliverable(
        sandbox_id=sid,
        file_rel=rel,
        task_spec=build_task_spec(
            task_type="multi_fact",
            headers=["用户ID", "总充值金额", "注册渠道"],
            time_window={"cohort": "pay"},
        ).to_dict(),
        column_plan=plan,
        trace=tr.to_dict(),
    )
    assert result["passed"], result.get("errors")
    assert result["ok"]
    assert result["status"] == "pass"
    assert result["summary"]["row_count"] >= 1
    assert result["summary"]["status"] == "pass"
    md = build_final_summary_from_verifier(result, time_label="vr")
    assert "总用户数" in md
    assert "### 导出概况" in md
    assert "校验状态" in md


def test_export_finalizer_updates_trace_repair_and_outcome(tmp_path, monkeypatch):
    sid = "sbx_finalizer"
    root = tmp_path / sid / "workplace"
    root.mkdir(parents=True)

    def _wp(s):
        return root if s == sid else Path("/nope")

    monkeypatch.setattr("app.services.workplace.ensure_workplace", _wp)
    monkeypatch.setattr("app.services.export_build_report.ensure_workplace", _wp)
    monkeypatch.setattr("app.services.export_verifier.ensure_workplace", _wp)
    monkeypatch.setattr("app.services.export_trace.ensure_workplace", _wp)

    run_id = "run_finalizer"
    write_task_json_page(
        sid,
        [{"uid": 1, "register_time": 1500, "sc": 10000}],
        run_id=run_id,
        page=1,
        view="view_result_user_info",
    )
    plan = build_column_plan(["用户ID"])
    rel = write_export_deliverable(
        sid,
        run_id,
        column_plan=plan,
        time_window={"start_ms": 1000, "end_ms": 2000, "label": "fin"},
        title="finalizer",
    )
    assert rel
    trace = ExportTrace(run_id=run_id)
    result = verify_export_finalizer(
        sandbox_id=sid,
        run_id=run_id,
        file_rel=rel,
        task_spec=build_task_spec(
            task_type="multi_fact",
            headers=["用户ID"],
            time_window={"start_ms": 1000, "end_ms": 2000},
        ).to_dict(),
        column_plan=plan,
        trace=trace,
    )
    assert result.verifier_result["status"] == "pass"
    assert result.repair_plan["status"] == "none"
    assert result.outcome.can_claim_complete
    loaded = read_export_trace(sid, run_id)
    assert loaded and loaded["verification"]["status"] == "pass"
    assert loaded["repair_plan"]["status"] == "none"


def test_verifier_blocks_core_failure_claim():
    tr = ExportTrace(run_id="x")
    tr.record_query(
        key="user", role="user", view="view_result_user_info",
        error="remote fail", status="error",
    )
    # No file → fail
    result = verify_export_deliverable(
        sandbox_id="missing",
        file_rel="nope.xlsx",
        task_spec={},
        column_plan=[{"header": "用户ID", "failure_policy": "core", "agg_spec": {"role": "user"}}],
        trace=tr.to_dict(),
        abandoned_roles={"user"},
    )
    assert not result["passed"]
    assert not result["ok"]
    assert result["status"] == "failed"
    assert "user" in result["missing_core"]
    assert result["errors"]


def test_verifier_marks_noncore_gap_repairable(tmp_path, monkeypatch):
    sid = "sbx_vr_repairable"
    root = tmp_path / sid / "workplace"
    root.mkdir(parents=True)

    def _wp(s):
        return root if s == sid else Path("/nope")

    monkeypatch.setattr("app.services.workplace.ensure_workplace", _wp)
    monkeypatch.setattr("app.services.export_build_report.ensure_workplace", _wp)
    monkeypatch.setattr("app.services.export_verifier.ensure_workplace", _wp)

    run_id = "run_vr_repairable"
    write_task_json_page(
        sid,
        [{"uid": 1, "register_time": 1500, "sc": 10000}],
        run_id=run_id,
        page=1,
        view="view_result_user_info",
    )
    plan = build_column_plan(["用户ID", "总下注金额(SC)"])
    rel = write_export_deliverable(
        sid,
        run_id,
        column_plan=plan,
        time_window={"start_ms": 1000, "end_ms": 2000, "label": "vr"},
        title="verifier",
    )
    result = verify_export_deliverable(
        sandbox_id=sid,
        file_rel=rel,
        task_spec=build_task_spec(
            task_type="multi_fact",
            headers=["用户ID", "总下注金额(SC)"],
            time_window={"start_ms": 1000, "end_ms": 2000},
        ).to_dict(),
        column_plan=plan,
        failed_views={"view_result_gameuser_betstat_everyday_bygame"},
    )
    assert result["passed"], result.get("errors")
    assert not result["ok"]
    assert result["status"] == "repairable"
    assert "总下注金额(SC)" in result["incomplete_non_core"]
    block = format_verifier_block(result)
    assert "状态: repairable" in block
    assert "修复建议" in block


def test_verifier_marks_zero_sequence_metric_repairable(tmp_path, monkeypatch):
    sid = "sbx_vr_zero_sequence"
    root = tmp_path / sid / "workplace"
    root.mkdir(parents=True)

    def _wp(s):
        return root if s == sid else Path("/nope")

    monkeypatch.setattr("app.services.export_verifier.ensure_workplace", _wp)

    import openpyxl

    dest = root / "zero_sequence.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["用户ID", "总充值金额", "连续充值次数(两次游戏行为之间的充值次数，需要最多的次数)"])
    ws.append(["7", 20, 0])
    ws.append(["8", 10, 0])
    notes = wb.create_sheet("口径说明")
    notes.append(["列名", "统计方法"])
    notes.append(["连续充值次数", "两次游戏行为之间的最大连续充值次数"])
    wb.save(dest)
    wb.close()

    plan = build_column_plan(["用户ID", "总充值金额", "连续充值次数"])
    result = verify_export_deliverable(
        sandbox_id=sid,
        file_rel="zero_sequence.xlsx",
        task_spec=build_task_spec(
            task_type="multi_fact",
            headers=["用户ID", "总充值金额", "连续充值次数"],
            time_window={"start_ms": 1000, "end_ms": 2000},
        ).to_dict(),
        column_plan=plan,
    )
    assert result["passed"], result.get("errors")
    assert not result["ok"]
    assert result["status"] == "repairable"
    assert any("连续充值次数" in c for c in result["incomplete_non_core"])
    assert any("指标列全 0" in w for w in result["warnings"])


def test_verifier_marks_low_cohort_rows_repairable(tmp_path, monkeypatch):
    sid = "sbx_vr_low_rows"
    root = tmp_path / sid / "workplace"
    root.mkdir(parents=True)

    def _wp(s):
        return root if s == sid else Path("/nope")

    monkeypatch.setattr("app.services.export_verifier.ensure_workplace", _wp)

    import openpyxl

    dest = root / "low_rows.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["用户ID", "总充值金额"])
    for i in range(10):
        ws.append([str(i + 1), 10])
    wb.create_sheet("口径说明").append(["列名", "统计方法"])
    wb.save(dest)
    wb.close()

    result = verify_export_deliverable(
        sandbox_id=sid,
        file_rel="low_rows.xlsx",
        task_spec=build_task_spec(
            task_type="multi_fact",
            headers=["用户ID", "总充值金额"],
            time_window={"start_ms": 1000, "end_ms": 2000},
        ).to_dict(),
        column_plan=build_column_plan(["用户ID", "总充值金额"]),
        trace={"export_contract": {"cohort_uid_estimate": 100}},
    )
    assert result["passed"], result.get("errors")
    assert not result["ok"]
    assert result["status"] == "repairable"
    assert "数据行数" in result["incomplete_non_core"]
    assert any("低于 cohort 目标" in w for w in result["warnings"])


def test_repair_plan_from_core_failure_uses_query_graph():
    plan = build_column_plan(["用户ID", "总充值金额"])
    graph = build_query_graph(plan, {"start_ms": 1, "end_ms": 2, "cohort": "pay"})
    trace = ExportTrace(run_id="repair")
    trace.column_plan = plan
    trace.query_graph = graph
    trace.done_node_keys = ["agg:pay:pay_sum"]
    verifier = {
        "status": "failed",
        "missing_core": ["user"],
        "missing_columns": [],
        "incomplete_non_core": [],
        "repair_hints": ["先补拉核心 role：user"],
    }
    repair = build_repair_plan(verifier, trace.to_dict())
    assert repair["status"] == "blocked"
    assert repair["source_status"] == "failed"
    action = repair["actions"][0]
    assert action["action_type"] == "fetch_core_role"
    assert action["role"] == "user"
    assert action["blocking"] is True
    assert "agg:pay:pay_sum" in repair["preserve_done_node_keys"]


def test_repair_plan_from_noncore_incomplete_skips_failed_node_once():
    plan = build_column_plan(["用户ID", "SC投注金额最多的游戏"])
    graph = build_query_graph(plan, {"start_ms": 1, "end_ms": 2})
    top_node = next(n for n in graph if n.get("mode") == "top_n")
    trace = ExportTrace(run_id="repair_noncore")
    trace.column_plan = plan
    trace.query_graph = graph
    trace.failed_views = [f"node:{top_node['key']}"]
    verifier = {
        "status": "repairable",
        "missing_core": [],
        "missing_columns": [],
        "incomplete_non_core": ["SC投注金额最多的游戏"],
        "repair_hints": ["补拉非核心缺口，或在口径说明中保留未完整标记"],
    }
    repair = build_repair_plan(verifier, trace.to_dict())
    assert repair["status"] == "repairable"
    assert repair["skip_failed_node_keys"] == [top_node["key"]]
    action = repair["actions"][0]
    assert action["action_type"] == "fetch_noncore_or_keep_incomplete"
    assert action["blocking"] is False
    assert top_node["key"] not in action["node_keys"]


def test_activate_repair_plan_from_trace_merges_resume_keys():
    activated = activate_repair_plan_from_trace({
        "done_node_keys": ["agg:pay"],
        "failed_views": ["view_result_old"],
        "repair_plan": {
            "status": "repairable",
            "preserve_done_node_keys": ["agg:pay", "agg:user"],
            "skip_failed_node_keys": ["top_n:bet:game"],
            "actions": [],
        },
    })
    assert activated["done_node_keys"] == ["agg:pay", "agg:user"]
    assert activated["failed_views"] == [
        "view_result_old",
        "node:top_n:bet:game",
    ]
    assert activated["repair_plan"]["status"] == "repairable"


def test_activate_repair_plan_from_prior_falls_back_to_run_state_repair_plan():
    activated = activate_repair_plan_from_prior(
        {
            "done_node_keys": ["agg:pay"],
            "failed_views": ["view_result_old"],
        },
        {
            "repair_plan": {
                "status": "repairable",
                "actions": [
                    {
                        "action_type": "describe_ads_views",
                        "views": ["ads_user_stats"],
                    },
                ],
                "preserve_done_node_keys": ["agg:user"],
            },
        },
    )
    assert activated["done_node_keys"] == ["agg:pay", "agg:user"]
    assert activated["failed_views"] == ["view_result_old"]
    assert activated["repair_plan"]["actions"][0]["action_type"] == "describe_ads_views"


def test_activate_repair_plan_from_prior_prefers_trace_repair_plan():
    activated = activate_repair_plan_from_prior(
        {
            "repair_plan": {
                "status": "blocked",
                "actions": [{"action_type": "fetch_core_role", "role": "pay"}],
            },
        },
        {
            "repair_plan": {
                "status": "repairable",
                "actions": [{"action_type": "describe_ads_views"}],
            },
        },
    )
    assert activated["repair_plan"]["status"] == "blocked"
    assert activated["repair_plan"]["actions"][0]["action_type"] == "fetch_core_role"


def test_prioritize_query_nodes_for_repair_moves_action_nodes_first():
    plan = build_column_plan(["用户ID", "总充值金额", "总提现金额"])
    nodes = build_query_graph(plan, {"start_ms": 1, "end_ms": 2, "cohort": "pay"})
    cash_node = next(n for n in nodes if n.get("role") == "cash")
    repair = {
        "status": "blocked",
        "actions": [
            {
                "action_type": "fetch_core_role",
                "priority": 0,
                "role": "cash",
                "views": [cash_node["view"]],
                "node_keys": [cash_node["key"]],
            }
        ],
    }
    ordered = prioritize_query_nodes_for_repair(nodes, repair)
    assert ordered[0]["key"] == cash_node["key"]
    assert {n["key"] for n in ordered} == {n["key"] for n in nodes}


def test_limit_query_nodes_for_repair_scopes_explicit_node_keys_only():
    plan = build_column_plan(["用户ID", "总充值金额", "总提现金额"])
    nodes = build_query_graph(plan, {"start_ms": 1, "end_ms": 2, "cohort": "pay"})
    cash_node = next(n for n in nodes if n.get("role") == "cash")
    scoped, did_scope = limit_query_nodes_for_repair(
        nodes,
        {
            "status": "blocked",
            "actions": [
                {
                    "action_type": "fetch_core_role",
                    "priority": 0,
                    "role": "cash",
                    "node_keys": [cash_node["key"]],
                }
            ],
        },
    )
    assert did_scope
    assert [n["key"] for n in scoped] == [cash_node["key"]]

    unscoped, did_scope = limit_query_nodes_for_repair(
        nodes,
        {
            "status": "repairable",
            "actions": [{"action_type": "inspect_verifier_warnings"}],
        },
    )
    assert not did_scope
    assert [n["key"] for n in unscoped] == [n["key"] for n in nodes]


def test_repair_plan_should_reverify_after_progress():
    repair = {
        "actions": [
            {
                "action_type": "fetch_noncore_or_keep_incomplete",
                "node_keys": ["top_n:bet:game"],
            }
        ]
    }
    assert repair_plan_should_reverify_after_progress(
        repair,
        landed_count=1,
        scoped=True,
    )
    assert repair_plan_should_reverify_after_progress(
        repair,
        landed_count=1,
        scoped=False,
    )
    assert not repair_plan_should_reverify_after_progress(
        repair,
        landed_count=0,
        scoped=True,
    )
    assert not repair_plan_should_reverify_after_progress(
        {"actions": [{"action_type": "inspect_verifier_warnings"}]},
        landed_count=1,
        scoped=False,
    )


def test_engine_outcome_gates_final_claims():
    passed = build_engine_outcome({"status": "pass", "summary": {"row_count": 1}})
    assert passed.can_claim_complete
    assert passed.can_deliver
    assert not passed.should_block_final

    repairable = build_engine_outcome(
        {
            "status": "repairable",
            "summary": {"row_count": 1},
            "repair_hints": ["补拉非核心缺口"],
        },
        {"status": "repairable"},
    )
    assert not repairable.can_claim_complete
    assert repairable.can_deliver
    assert not repairable.should_block_final
    assert repairable.repair_hints == ["补拉非核心缺口"]

    failed = build_engine_outcome(
        {
            "status": "failed",
            "blocking_reasons": ["核心缺失: user"],
        },
        {"status": "blocked"},
    )
    assert not failed.can_claim_complete
    assert not failed.can_deliver
    assert failed.should_block_final
    assert failed.blocking_reasons == ["核心缺失: user"]

    fallback = build_engine_outcome({"status": "failed"}, mode="fallback")
    assert not fallback.should_block_final


def test_soft_fail_sequence_sql_still_builds():
    plan = build_column_plan(["连续充值次数"])
    view, sql = sql_from_agg_spec(plan[0]["agg_spec"], {"start_ms": 1, "end_ms": 2})
    assert view == "view_result_pay_order_log"
    assert "create_time" in sql
    nodes = build_query_graph(
        build_column_plan(["用户ID", "总充值金额", "连续充值次数"]),
        {"start_ms": 1, "end_ms": 2, "cohort": "pay"},
    )
    assert any(n.get("mode") == "sequence" for n in nodes)
